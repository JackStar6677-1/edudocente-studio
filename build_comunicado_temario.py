import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_section_header,
    set_cell_shading, set_cell_borders, set_cell_margins, COLOR_NAVY_HEX,
    COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_NAVY_RGB,
    COLOR_DARK_RGB, COLOR_WHITE_RGB, OUTPUT_DIR
)

def build_comunicado():
    doc = create_base_doc()

    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="INFORMATIVO PEDAGÓGICO - 5° BÁSICO A Y B")
    add_title_banner(doc, "TEMARIO Y CALENDARIZACIÓN EVALUACIÓN FINAL DE CIENCIAS NATURALES")

    # CUADRO DE DATOS GENERALES
    t_info = doc.add_table(rows=2, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False

    w_left, w_right = Pt(255), Pt(256.2)
    datos = [
        ("Estimada Comunidad Escolar (5° A y B):", "Asignatura: Ciencias Naturales"),
        ("Objetivo de Aprendizaje: CN05 OA 11", "Puntaje Total: 28 puntos")
    ]

    for r_idx, (d1, d2) in enumerate(datos):
        c1, c2 = t_info.cell(r_idx, 0), t_info.cell(r_idx, 1)
        c1.width, c2.width = w_left, w_right
        for c, text in [(c1, d1), (c2, d2)]:
            set_cell_margins(c, top=60, bottom=60, left=90, right=90)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            if r_idx == 0:
                set_cell_shading(c, "F6FAFE")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if "Estimada" in text or "Asignatura" in text:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # MENSAJE DE PRESENTACIÓN
    p_saludo = doc.add_paragraph()
    p_saludo.paragraph_format.space_before = Pt(2)
    p_saludo.paragraph_format.space_after = Pt(4)
    r_sal = p_saludo.add_run(
        "Junto con saludar cordialmente a nuestros estudiantes, madres, padres y apoderados, se hace entrega del temario oficial "
        "correspondiente a la Evaluación Final de Ciencias Naturales para 5° Básico A y B. Invitamos a las familias a acompañar "
        "y reforzar el proceso de estudio en el hogar con base en las orientaciones pedagógicas detalladas a continuación:"
    )
    r_sal.font.name = "Arial"
    r_sal.font.size = Pt(8.5)

    # SECCIÓN 1: CONTENIDOS Y PÁGINAS DEL TEXTO
    add_section_header(doc, "1. CONTENIDOS Y REFERENCIAS DEL TEXTO ESCOLAR MINEDUC")

    t_temas = doc.add_table(rows=5, cols=3)
    t_temas.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_temas.autofit = False

    t_headers = ["Contenido Específico", "Páginas Texto Mineduc", "Conceptos Clave a Reforzar"]
    t_widths = [Pt(150), Pt(110), Pt(251.2)]

    for i, h in enumerate(t_headers):
        c = t_temas.cell(0, i)
        c.width = t_widths[i]
        set_cell_shading(c, COLOR_NAVY_HEX)
        set_cell_margins(c, top=60, bottom=60, left=70, right=70)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = COLOR_WHITE_RGB

    filas_temario = [
        ("¿Qué es la energía eléctrica?", "Página 146", "Movimiento de cargas eléctricas, partículas de la materia, cargas positivas (+) y negativas (-)."),
        ("Circuitos eléctricos", "Página 149", "Concepto de circuito cerrado, función de la fuente/generador, receptor (transformación de energía) e interruptor."),
        ("Símbolos de circuitos", "Página 150", "Representación estandarizada: símbolo de pila (+ y -), cables conductores, interruptor y ampolleta/resistencia."),
        ("Conductores eléctricos", "Página 152", "Materiales que facilitan el desplazamiento de cargas (ej. cobre) frente a materiales aislantes (plástico, madera, goma).")
    ]

    for r_i, (c1_val, c2_val, c3_val) in enumerate(filas_temario, start=1):
        for c_i, val in enumerate([c1_val, c2_val, c3_val]):
            c = t_temas.cell(r_i, c_i)
            c.width = t_widths[c_i]
            set_cell_margins(c, top=50, bottom=50, left=70, right=70)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            if c_i == 0:
                set_cell_shading(c, "F6FAFE")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if c_i <= 1:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # SECCIÓN 2: ESTRUCTURA DE LA PRUEBA
    add_section_header(doc, "2. ESTRUCTURA DE LA EVALUACIÓN (TOTAL: 28 PUNTOS)")

    items_desc = [
        ("• Ítem I: Selección Múltiple (10 puntos):", " 5 preguntas de alternativas sobre conceptos de energía eléctrica, materia y conductores (2 puntos cada una)."),
        ("• Ítem II: Verdadero o Falso (10 puntos):", " 10 afirmaciones sobre circuitos eléctricos, componentes y funcionamiento del interruptor (1 punto cada una)."),
        ("• Ítem III: Dibujo de Símbolos (8 puntos):", " Dibujo de los símbolos esquemáticos de 4 componentes: fuente o generador, cables, interruptor y resistencia (2 puntos cada uno).")
    ]

    for it_title, it_desc in items_desc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        rt = p.add_run(it_title)
        rt.bold = True
        rt.font.name = "Arial"
        rt.font.size = Pt(8.5)
        rt.font.color.rgb = COLOR_NAVY_RGB
        rd = p.add_run(it_desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # SECCIÓN 3: RECOMENDACIONES DE ESTUDIO Y MATERIALES
    add_section_header(doc, "3. RECOMENDACIONES Y SUGERENCIAS DE ESTUDIO")

    consejos = [
        ("Materiales sugeridos: ", "Se recuerda traer lápiz grafito o pasta (azul o negro), goma de borrar y regla para trazar adecuadamente los símbolos de circuitos."),
        ("Práctica de símbolos: ", "Se sugiere practicar en casa el trazado de los 4 símbolos fundamentales: fuente o generador (+ y -), cables, interruptor y resistencia/ampolleta."),
        ("Repaso del libro: ", "Recuerda revisar y leer comprensivamente las páginas 146, 149, 150 y 152 del texto escolar de Ciencias Naturales entregado por el MINEDUC."),
        ("Consultas y dudas: ", "Opcionalmente, los estudiantes pueden acercarse a resolver inquietudes previas a la evaluación durante las horas de clase.")
    ]

    for c_title, c_text in consejos:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2)
        r_num = p.add_run("✓ ")
        r_num.bold = True
        r_num.font.color.rgb = COLOR_NAVY_RGB
        r_t = p.add_run(c_title)
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.5)
        r_txt = p.add_run(c_text)
        r_txt.font.name = "Arial"
        r_txt.font.size = Pt(8.5)

    # Contacto y Firma de la Profesora Margarita
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(18)
    p_sign.paragraph_format.space_after = Pt(0)
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_mail = p_sign.add_run("Correo institucional de contacto: profesora.margaritamiranda@cepluispasteur.cl\n\n")
    r_mail.font.name = "Arial"
    r_mail.font.size = Pt(8.5)
    r_mail.italic = True
    r_mail.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    rs1 = p_sign.add_run("Saludos cordiales\n")
    rs1.font.name = "Arial"
    rs1.font.size = Pt(9)
    rs1.font.color.rgb = COLOR_DARK_RGB

    rs2 = p_sign.add_run("Profesora Margarita Miranda B.\n")
    rs2.bold = True
    rs2.font.name = "Arial"
    rs2.font.size = Pt(9.5)
    rs2.font.color.rgb = COLOR_NAVY_RGB

    rs3 = p_sign.add_run("Profesora de Educación Básica\nC.E.P. Luis Pasteur Anexo")
    rs3.font.name = "Arial"
    rs3.font.size = Pt(8.5)
    rs3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    out_path = os.path.join(OUTPUT_DIR, "Temario_Evaluacion_Final_Ciencias_5Basico.docx")
    doc.save(out_path)
    print(f"Temario actualizado con éxito en: {out_path}")

if __name__ == "__main__":
    build_comunicado()
