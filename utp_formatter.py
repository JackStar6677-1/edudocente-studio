#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de Procesamiento y Estandarización de Evaluaciones Finales 2026
Directrices de UTP:
1. Formato: Letra Arial tamaño 12 (ajustable en tablas para legibilidad institucional).
2. Puntaje total: Mínimo 25 puntos.
3. Estructura: Al menos 3 ítems variados con puntaje explícito por dificultad.
4. Separación estricta:
   - Prueba del Estudiante: 100% limpia sin respuestas ni pautas.
   - Pauta de Corrección Docente: Documento anexo independiente con solucionario detallado y Tabla de Especificaciones oficial (o tabla de rúbrica).
   - Temarios por Curso: Documento oficial para apoderados y estudiantes (plazo 9 de octubre).
"""

import os
import sys
import shutil
import zipfile
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Constantes de estilo institucional
COLOR_NAVY_HEX = "173F73"
COLOR_ICE_HEX = "EAF3FB"
COLOR_BOX_BG_HEX = "F6FAFE"
COLOR_BORDER_HEX = "B0C4DE"

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, **kwargs):
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            sz = edge_data.get('sz', '4')
            color = edge_data.get('color', COLOR_BORDER_HEX)
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    cell._tc.get_or_add_tcPr().append(tcBorders)

def standardize_document_font(doc, default_size=12):
    """
    Recorre todo el documento asegurando Arial 12 para texto regular,
    manteniendo títulos y membretes en jerarquía visual armónica.
    """
    for p in doc.paragraphs:
        for run in p.runs:
            run.font.name = 'Arial'
            if run.font.size is None:
                run.font.size = Pt(default_size)
            elif run.font.size.pt < 11.5 and not run.bold:
                run.font.size = Pt(default_size)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = 'Arial'
                        if run.font.size is None:
                            run.font.size = Pt(11)
                        elif run.font.size.pt < 10:
                            run.font.size = Pt(10.5)

def append_pauta_summary_table(doc, items_data):
    """
    Agrega la tabla resumen de pauta de corrección con formato institucional:
    ITEM | PREGUNTA | RESPUESTA | PUNTAJE
    """
    doc.add_page_break()
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(6)
    p_header.paragraph_format.space_after = Pt(8)
    r = p_header.add_run("TABLA DE CLAVES Y RESPUESTAS (PAUTA OFICIAL)")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0x17, 0x3F, 0x73)

    table = doc.add_table(rows=len(items_data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Pt(60), Pt(80), Pt(280), Pt(80)]
    headers = ["ÍTEM", "N°", "RESPUESTA CORRECTA / CRITERIO", "PUNTAJE"]

    for col_idx, (h_text, w) in enumerate(zip(headers, widths)):
        cell = table.cell(0, col_idx)
        cell.width = w
        set_cell_shading(cell, COLOR_NAVY_HEX)
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, data in enumerate(items_data, start=1):
        item_lbl, q_num, ans_text, pts_text = data
        bg_col = "FFFFFF" if row_idx % 2 != 0 else COLOR_BOX_BG_HEX
        for col_idx, (val, w) in enumerate(zip([item_lbl, q_num, ans_text, pts_text], widths)):
            cell = table.cell(row_idx, col_idx)
            cell.width = w
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            set_cell_borders(cell, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                             left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if col_idx in (0, 1, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(10)
            if col_idx == 3:
                run.bold = True

def append_specification_table(doc, table_data):
    """
    Agrega la Tabla de Especificaciones según formato UTP:
    Ejes / Contenidos / Ítems | Habilidades cognitivas con porcentajes | Total
    """
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(18)
    p_header.paragraph_format.space_after = Pt(8)
    r = p_header.add_run("TABLA DE ESPECIFICACIONES (UNIDAD TÉCNICA PEDAGÓGICA)")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = RGBColor(0x17, 0x3F, 0x73)

    headers = table_data["headers"]
    rows = table_data["rows"]
    widths = table_data.get("widths", [Pt(150), Pt(100), Pt(100), Pt(100), Pt(70)])

    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    for col_idx, (h_text, w) in enumerate(zip(headers, widths)):
        cell = tbl.cell(0, col_idx)
        cell.width = w
        set_cell_shading(cell, COLOR_NAVY_HEX)
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, r_data in enumerate(rows, start=1):
        bg_col = "FFFFFF" if r_idx % 2 != 0 else COLOR_BOX_BG_HEX
        is_total_row = (r_idx == len(rows))
        if is_total_row:
            bg_col = "E2F0D9"

        for c_idx, (val, w) in enumerate(zip(r_data, widths)):
            cell = tbl.cell(r_idx, c_idx)
            cell.width = w
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            set_cell_borders(cell, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                             left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(10)
            if is_total_row or c_idx == len(r_data) - 1:
                run.bold = True

if __name__ == '__main__':
    print("Módulo utp_formatter listo.")
