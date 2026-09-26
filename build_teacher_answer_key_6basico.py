import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from generator_core import (
    create_base_doc, add_header, add_title_banner,
    add_section_header, set_cell_shading, set_cell_borders, set_cell_margins,
    COLOR_NAVY_HEX, COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX,
    COLOR_CORRECT_HEX, COLOR_NAVY_RGB, COLOR_CORRECT_RGB, COLOR_DARK_RGB,
    COLOR_WHITE_RGB, OUTPUT_DIR, MATTER_ASSETS_DIR
)
from build_student_test_6basico import add_student_info_6b

def build_pauta_correccion_6b():
    doc = create_base_doc()

    # ================== PÁGINA 1 ==================
    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="CIENCIAS NATURALES - 6° BÁSICO A - B", subtitle="DOCUMENTO DOCENTE")
    add_title_banner(doc, "PAUTA DE CORRECCIÓN: EVALUACIÓN FINAL DE CIENCIAS NATURALES", is_pauta=True)
    add_student_info_6b(doc, is_pauta=True)

    # SOLUCIONARIO ÍTEM I
    add_section_header(doc, "SOLUCIONARIO ÍTEM I: SELECCIÓN MÚLTIPLE (2 pts c/u — Total: 10 pts)")

    pautas_item1 = [
        ("1. Los cambios de estado de la materia son:",
         "C) Los cambios físicos en los que la materia cambia su aspecto externo, sin alterar su composición.",
         "Justificación: En los cambios de estado las sustancias experimentan variaciones en su forma física, volumen y orden de partículas, pero sus moléculas continúan siendo las mismas (Texto escolar Pág. 170)."),

        ("2. En un cambio de estado, la materia solo cambia su aspecto...",
         "A) pero continúa siendo la misma sustancia.",
         "Justificación: El agua líquida, el hielo y el vapor de agua siguen siendo agua ($H_2O$), no se transforma en una sustancia distinta (Texto escolar Pág. 170)."),

        ("3. Los cambios de estado se producen por:",
         "B) Por absorción o por liberación de energía (calor).",
         "Justificación: Al suministrar calor (absorción) las partículas ganan energía cinética y se separan; al enfriar (liberación) pierden energía y se acercan (Págs. 173 y 181)."),

        ("4. Los cambios progresivos de la materia son:",
         "A) Los cambios por absorción de energía (calor).",
         "Justificación: Los cambios progresivos requieren incorporar calor del entorno: fusión, vaporización y sublimación progresiva (Texto escolar Pág. 173)."),

        ("5. Los cambios regresivos de la materia son aquellos que ocurren:",
         "A) Por liberación de calor (enfriamiento).",
         "Justificación: Los cambios regresivos ceden o liberan calor al ambiente al enfriarse: condensación, solidificación y sublimación regresiva/inversa (Texto escolar Pág. 181).")
    ]

    for q_t, ans, just in pautas_item1:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(2)
        p_q.paragraph_format.space_after = Pt(1)
        rq = p_q.add_run(q_t)
        rq.bold = True
        rq.font.name = "Arial"
        rq.font.size = Pt(8.5)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Inches(0.2)
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(0.5)
        ra_b = p_ans.add_run("Respuesta Correcta: ")
        ra_b.bold = True
        ra_b.font.name = "Arial"
        ra_b.font.size = Pt(8.5)
        ra_b.font.color.rgb = COLOR_CORRECT_RGB
        ra_t = p_ans.add_run(ans)
        ra_t.bold = True
        ra_t.font.name = "Arial"
        ra_t.font.size = Pt(8.5)
        ra_t.font.color.rgb = COLOR_CORRECT_RGB

        p_just = doc.add_paragraph()
        p_just.paragraph_format.left_indent = Inches(0.2)
        p_just.paragraph_format.space_before = Pt(0)
        p_just.paragraph_format.space_after = Pt(2)
        rj = p_just.add_run(just)
        rj.font.name = "Arial"
        rj.font.size = Pt(8)
        rj.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_page_break()

    # ================== PÁGINA 2 ==================
    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="CIENCIAS NATURALES - 6° BÁSICO A - B", subtitle="PAUTA DE CORRECCIÓN (PÁG. 2)")

    # SOLUCIONARIO ÍTEM II
    add_section_header(doc, "SOLUCIONARIO ÍTEM II: VERDADERO O FALSO (1 pto c/u — Total: 10 pts)")

    pautas_vf = [
        ("V", "1. Los cambios de estado por absorción de calor son fusión, vaporización y sublimación.",
         "Verdadero: En estos tres cambios las partículas absorben energía térmica para aumentar su separación y movimiento."),
        
        ("V", "2. El proceso en el cual un líquido pasa a estado gaseoso es la vaporización.",
         "Verdadero: La vaporización es el concepto general que engloba el paso de estado líquido a gaseoso."),
        
        ("F", "3. La fusión es el proceso de un cuerpo en estado sólido a gaseoso.",
         "Falsa: La fusión es el paso de SÓLIDO A LÍQUIDO. El paso de sólido a gas se llama sublimación progresiva."),
        
        ("V", "4. El proceso de vaporización puede ocurrir de dos formas: evaporación y ebullición.",
         "Verdadero: La evaporación ocurre solo en la superficie y a cualquier temperatura; la ebullición ocurre en toda la masa al alcanzar el punto de ebullición."),
        
        ("F", "5. En la ebullición participan pocas o algunas partículas de la superficie del líquido.",
         "Falsa: En la ebullición participan TODAS las partículas del líquido (formando burbujas en todo el volumen). La evaporación es la que ocurre solo en la superficie."),
        
        ("V", "6. Algunos ejemplos de sustancias que experimentan sublimación son el yodo y la naftalina.",
         "Verdadero: La naftalina y el yodo pasan de sólido a gas directamente a temperatura ambiente."),
        
        ("V", "7. Los tres estados de la materia estudiados son sólido, líquido y gaseoso.",
         "Verdadero: Son los tres estados fundamentales analizados en el currículum de 6° básico."),
        
        ("F", "8. Los únicos cambios de estado por liberación de calor son la solidificación y la condensación.",
         "Falsa: Falta la SUBLIMACIÓN INVERSA (o regresiva), que también se produce por liberación de calor."),
        
        ("V", "9. La solidificación es el proceso en que un líquido pasa a estado sólido.",
         "Verdadero: Al enfriarse, el líquido pierde energía térmica y sus partículas se ordenan en estado sólido."),
        
        ("V", "10. La condensación es el proceso de un gas cuando pasa a estado líquido.",
         "Verdadero: Al perder calor, las partículas gaseosas reducen su movimiento y se condensan en gotas de líquido.")
    ]

    for v_f, sent, just in pautas_vf:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)

        badge = f"[  {v_f}  ]  "
        rb = p.add_run(badge)
        rb.bold = True
        rb.font.name = "Arial"
        rb.font.size = Pt(8.5)
        if v_f == "V":
            rb.font.color.rgb = COLOR_CORRECT_RGB
        else:
            rb.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

        rs = p.add_run(sent)
        rs.bold = True
        rs.font.name = "Arial"
        rs.font.size = Pt(8.5)

        pj = doc.add_paragraph()
        pj.paragraph_format.left_indent = Inches(0.4)
        pj.paragraph_format.space_before = Pt(0)
        pj.paragraph_format.space_after = Pt(2.5)
        rj = pj.add_run(f"→ {just}")
        rj.font.name = "Arial"
        rj.font.size = Pt(8)
        rj.italic = True
        rj.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # SOLUCIONARIO ÍTEM III
    add_section_header(doc, "SOLUCIONARIO ÍTEM III: DIBUJOS DE CAMBIOS POR LIBERACIÓN DE CALOR (Total: 6 pts)")

    procesos_pauta = [
        {"nombre": "1. Solidificación (2 pts)",
         "ini_img": "estado_liquido.png", "fin_img": "estado_solido.png",
         "desc_ini": "Estado Inicial: Líquido\n(Partículas cercanas pero desordenadas)",
         "desc_fin": "Estado Final: Sólido\n(Partículas ordenadas y compactas)",
         "criterio": "2 pts: Representa el paso de líquido a sólido con partículas que pasan de desordenadas a compactas/ordenadas."},

        {"nombre": "2. Sublimación Inversa (2 pts)",
         "ini_img": "estado_gaseoso.png", "fin_img": "estado_solido.png",
         "desc_ini": "Estado Inicial: Gas / Vapor\n(Partículas muy separadas y dispersas)",
         "desc_fin": "Estado Final: Sólido\n(Partículas compactas y ordenadas)",
         "criterio": "2 pts: Representa el paso directo de gas a sólido sin pasar por el estado líquido intermedio."},

        {"nombre": "3. Condensación (2 pts)",
         "ini_img": "estado_gaseoso.png", "fin_img": "estado_liquido.png",
         "desc_ini": "Estado Inicial: Gas / Vapor\n(Partículas muy separadas y en movimiento)",
         "desc_fin": "Estado Final: Líquido\n(Partículas agrupadas en el fondo)",
         "criterio": "2 pts: Representa el paso de gas a líquido (o dibuja formación de gotitas de agua al enfriarse)."}
    ]

    for pr in procesos_pauta:
        t_proc = doc.add_table(rows=1, cols=3)
        t_proc.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_proc.autofit = False

        c_ini, c_arr, c_fin = t_proc.cell(0, 0), t_proc.cell(0, 1), t_proc.cell(0, 2)
        c_ini.width, c_arr.width, c_fin.width = Pt(210), Pt(91.2), Pt(210)

        for c in [c_ini, c_fin]:
            set_cell_margins(c, top=40, bottom=40, left=50, right=50)
            set_cell_borders(c, top={'val': 'single', 'sz': '6', 'color': COLOR_CORRECT_HEX},
                                bottom={'val': 'single', 'sz': '6', 'color': COLOR_CORRECT_HEX},
                                left={'val': 'single', 'sz': '6', 'color': COLOR_CORRECT_HEX},
                                right={'val': 'single', 'sz': '6', 'color': COLOR_CORRECT_HEX})
            set_cell_shading(c, "F0FFF4")

        set_cell_margins(c_arr, top=40, bottom=40, left=20, right=20)
        set_cell_borders(c_arr, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})

        # Inicial
        p_i = c_ini.paragraphs[0]
        p_i.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_i.paragraph_format.space_before = Pt(0)
        p_i.paragraph_format.space_after = Pt(2)
        r_tit = p_i.add_run(f"{pr['nombre']}\n")
        r_tit.bold = True
        r_tit.font.name = "Arial"
        r_tit.font.size = Pt(8.5)
        r_tit.font.color.rgb = COLOR_NAVY_RGB

        f_ini = os.path.join(MATTER_ASSETS_DIR, pr['ini_img'])
        if os.path.exists(f_ini):
            p_i.add_run().add_picture(f_ini, width=Inches(0.9))
        r_di = p_i.add_run(f"\n{pr['desc_ini']}")
        r_di.font.name = "Arial"
        r_di.font.size = Pt(7.5)

        # Flecha
        p_a = c_arr.paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_a.paragraph_format.space_before = Pt(12)
        p_a.paragraph_format.space_after = Pt(0)
        ra_t = p_a.add_run("Liberación\nde Calor\n\n")
        ra_t.bold = True
        ra_t.font.name = "Arial"
        ra_t.font.size = Pt(7.5)
        ra_t.font.color.rgb = COLOR_CORRECT_RGB
        ra_sym = p_a.add_run("────►")
        ra_sym.bold = True
        ra_sym.font.name = "Arial"
        ra_sym.font.size = Pt(12)
        ra_sym.font.color.rgb = COLOR_CORRECT_RGB

        # Final
        p_f = c_fin.paragraphs[0]
        p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_f.paragraph_format.space_before = Pt(0)
        p_f.paragraph_format.space_after = Pt(2)
        r_tit2 = p_f.add_run("Resultado:\n")
        r_tit2.bold = True
        r_tit2.font.name = "Arial"
        r_tit2.font.size = Pt(8.5)
        r_tit2.font.color.rgb = COLOR_NAVY_RGB

        f_fin = os.path.join(MATTER_ASSETS_DIR, pr['fin_img'])
        if os.path.exists(f_fin):
            p_f.add_run().add_picture(f_fin, width=Inches(0.9))
        r_df = p_f.add_run(f"\n{pr['desc_fin']}")
        r_df.font.name = "Arial"
        r_df.font.size = Pt(7.5)

        # Criterio
        p_cr = doc.add_paragraph()
        p_cr.paragraph_format.left_indent = Inches(0.15)
        p_cr.paragraph_format.space_before = Pt(2)
        p_cr.paragraph_format.space_after = Pt(3)
        rc = p_cr.add_run(f"Criterio: {pr['criterio']}")
        rc.font.name = "Arial"
        rc.font.size = Pt(7.5)
        rc.italic = True
        rc.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    out_path = os.path.join(OUTPUT_DIR, "Pauta_Correccion_Ciencias_6Basico.docx")
    doc.save(out_path)
    print(f"Pauta 6to Basico generated successfully: {out_path}")

if __name__ == "__main__":
    build_pauta_correccion_6b()
