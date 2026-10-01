import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_student_info,
    add_instructions_box, add_section_header, set_cell_shading,
    set_cell_borders, set_cell_margins, COLOR_NAVY_HEX, COLOR_ICE_HEX,
    COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_DARK_RGB, COLOR_NAVY_RGB,
    OUTPUT_DIR
)

def build_evaluacion_estudiante():
    doc = create_base_doc()

    # ================== PÁGINA 1 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 5° BÁSICO A - B")
    add_title_banner(doc, "EVALUACIÓN FINAL: ENERGÍA ELÉCTRICA Y CIRCUITOS")
    add_student_info(doc, is_pauta=False)
    add_instructions_box(doc)

    # ÍTEM I
    add_section_header(doc, "ÍTEM I: SELECCIÓN MÚLTIPLE (2 puntos c/u — Total: 10 puntos)")
    
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(1)
    p_inst.paragraph_format.space_after = Pt(4)
    r_inst = p_inst.add_run("Lee atentamente cada pregunta y marca con una ")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(8.5)
    r_x = p_inst.add_run("X")
    r_x.bold = True
    r_x.font.name = "Arial"
    r_x.font.size = Pt(8.5)
    r_inst2 = p_inst.add_run(" la alternativa que consideres correcta.")
    r_inst2.font.name = "Arial"
    r_inst2.font.size = Pt(8.5)

    preguntas_item1 = [
        ("1. ¿Qué es la energía eléctrica?",
         [("A", "Una forma de energía muy escasa que solo se produce en laboratorios nucleares."),
          ("B", "Una forma de energía producida exclusivamente por imanes naturales."),
          ("C", "Una forma de energía que se produce por el movimiento y flujo de cargas eléctricas."),
          ("D", "Un tipo de energía que produce el cuerpo humano para comunicarse.")]),

        ("2. Esta energía surge de las propiedades de la materia. Desde el punto de vista científico, ¿cómo se describe la estructura fundamental de la materia?",
         [("A", "La materia está formada por partículas diminutas que poseen cargas eléctricas."),
          ("B", "La materia está compuesta exclusivamente por redes de conexiones eléctricas."),
          ("C", "La materia es únicamente aquella sustancia sólida que compone la corteza terrestre."),
          ("D", "La materia está constituida por rayos luminosos en constante movimiento en el espacio.")]),

        ("3. En los modelos científicos que estudiamos en la escuela, las partículas que componen la materia se suelen representar como:",
         [("A", "Luces y destellos solares microscópicos."),
          ("B", "Pequeñas esferas."),
          ("C", "Cargas positivas únicamente."),
          ("D", "Cargas negativas únicamente.")]),

        ("4. Las partículas pueden adquirir o manifestar una carga eléctrica, la que puede ser:",
         [("A", "Solar o eólica."),
          ("B", "Positiva (+) o negativa (-)."),
          ("C", "Imanes y magnetos."),
          ("D", "Cinética - química.")]),

        ("5. Las cargas eléctricas negativas pueden desplazarse con facilidad a través de algunos materiales conductores. ¿Cuál de los siguientes materiales es un conductor eléctrico comúnmente utilizado en cables?",
         [("A", "Plástico."),
          ("B", "Cobre."),
          ("C", "Madera seca."),
          ("D", "Goma o caucho.")])
    ]

    for q_text, options in preguntas_item1:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        rq = p_q.add_run(q_text)
        rq.font.name = "Arial"
        rq.font.size = Pt(9)
        rq.bold = True
        rq.font.color.rgb = COLOR_DARK_RGB

        for letter, opt_text in options:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.2)
            p_opt.paragraph_format.space_before = Pt(0.5)
            p_opt.paragraph_format.space_after = Pt(1)
            
            # Solo A), B), C), D) sin corchetes
            rl = p_opt.add_run(f"{letter})  ")
            rl.font.name = "Arial"
            rl.font.size = Pt(8.5)
            rl.bold = True
            rl.font.color.rgb = COLOR_NAVY_RGB

            ro = p_opt.add_run(opt_text)
            ro.font.name = "Arial"
            ro.font.size = Pt(8.5)
            ro.font.color.rgb = COLOR_DARK_RGB

    # SALTO A PÁGINA 2
    doc.add_page_break()

    # ================== PÁGINA 2 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 5° BÁSICO A - B", subtitle="EVALUACIÓN FINAL (PÁG. 2)")
    
    # ÍTEM II
    add_section_header(doc, "ÍTEM II: VERDADERO O FALSO (1 punto c/u — Total: 10 puntos)")

    p_vf_inst = doc.add_paragraph()
    p_vf_inst.paragraph_format.space_before = Pt(1)
    p_vf_inst.paragraph_format.space_after = Pt(4)
    r_vf = p_vf_inst.add_run("Lee atentamente las siguientes afirmaciones. Escribe una ")
    r_vf.font.name = "Arial"
    r_vf.font.size = Pt(8.5)
    r_v = p_vf_inst.add_run("V")
    r_v.bold = True
    r_v.font.name = "Arial"
    r_v.font.size = Pt(8.5)
    r_vf2 = p_vf_inst.add_run(" si la afirmación es verdadera o una ")
    r_vf2.font.name = "Arial"
    r_vf2.font.size = Pt(8.5)
    r_f = p_vf_inst.add_run("F")
    r_f.bold = True
    r_f.font.name = "Arial"
    r_f.font.size = Pt(8.5)
    r_vf3 = p_vf_inst.add_run(" si es falsa.")
    r_vf3.font.name = "Arial"
    r_vf3.font.size = Pt(8.5)

    afirmaciones = [
        "1.  ( _____ )  Los aparatos electrónicos y electrodomésticos funcionan gracias a los circuitos eléctricos.",
        "2.  ( _____ )  Un circuito eléctrico es un camino cerrado por el que fluyen y circulan las cargas eléctricas.",
        "3.  ( _____ )  Todos los circuitos eléctricos tienen componentes que funcionan sin conectarse entre sí.",
        "4.  ( _____ )  El generador o fuente de energía proporciona la energía necesaria para que circule la corriente eléctrica.",
        "5.  ( _____ )  Los cables conducen la corriente eléctrica por las paredes de la casa de forma espontánea y sin fuente de energía.",
        "6.  ( _____ )  La resistencia o receptor recibe y transforma la energía eléctrica en otro tipo de energía.",
        "7.  ( _____ )  Si el receptor es una ampolleta, la energía eléctrica se transforma principalmente en energía lumínica.",
        "8.  ( _____ )  El interruptor regula el paso de la corriente hidráulica.",
        "9.  ( _____ )  Cuando el interruptor está abierto, la corriente eléctrica deja de circular por el circuito.",
        "10. ( _____ )  Cuando el interruptor está cerrado, la corriente eléctrica circula con normalidad por el circuito."
    ]

    for af in afirmaciones:
        p_af = doc.add_paragraph()
        p_af.paragraph_format.left_indent = Inches(0.15)
        p_af.paragraph_format.space_before = Pt(1.5)
        p_af.paragraph_format.space_after = Pt(2.5)
        r = p_af.add_run(af)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_DARK_RGB

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(2)

    # ÍTEM III
    add_section_header(doc, "ÍTEM III: DIBUJO DE SÍMBOLOS SEGÚN CORRESPONDA (2 puntos c/u — Total: 8 puntos)")

    p_i3_inst = doc.add_paragraph()
    p_i3_inst.paragraph_format.space_before = Pt(1)
    p_i3_inst.paragraph_format.space_after = Pt(4)
    r_i3 = p_i3_inst.add_run("Los circuitos eléctricos se representan mediante esquemas utilizando símbolos. ")
    r_i3.font.name = "Arial"
    r_i3.font.size = Pt(8.5)
    r_i3_b = p_i3_inst.add_run("Dibuja los símbolos según corresponda en cada cuadro:")
    r_i3_b.bold = True
    r_i3_b.font.name = "Arial"
    r_i3_b.font.size = Pt(8.5)

    # Tabla 2x2: SOLO los 4 nombres solicitados, sin pistas ni pistas entre paréntesis
    t_cuadros = doc.add_table(rows=2, cols=2)
    t_cuadros.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cuadros.autofit = False

    componentes = [
        {"title": "Fuente o generador", "row": 0, "col": 0},
        {"title": "Cables", "row": 0, "col": 1},
        {"title": "Interruptor", "row": 1, "col": 0},
        {"title": "Resistencia", "row": 1, "col": 1},
    ]

    for comp in componentes:
        cell = t_cuadros.cell(comp["row"], comp["col"])
        cell.width = Pt(250)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_borders(cell, top={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               bottom={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               left={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               right={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX})
        set_cell_shading(cell, "FAFCFF")

        # Título del cuadro: ÚNICAMENTE el nombre formal, sin pistas
        p_head = cell.paragraphs[0]
        p_head.paragraph_format.space_before = Pt(0)
        p_head.paragraph_format.space_after = Pt(2)
        r_th = p_head.add_run(comp["title"])
        r_th.bold = True
        r_th.font.name = "Arial"
        r_th.font.size = Pt(9.5)
        r_th.font.color.rgb = COLOR_NAVY_RGB

        # Recuadro interior limpio para que el estudiante dibuje
        t_draw = cell.add_table(rows=1, cols=1)
        t_draw.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_draw.autofit = False
        c_draw = t_draw.cell(0, 0)
        c_draw.width = Pt(230)
        set_cell_margins(c_draw, top=80, bottom=80, left=80, right=80)
        set_cell_borders(c_draw, top={'val': 'dashed', 'sz': '6', 'color': '999999'},
                                 bottom={'val': 'dashed', 'sz': '6', 'color': '999999'},
                                 left={'val': 'dashed', 'sz': '6', 'color': '999999'},
                                 right={'val': 'dashed', 'sz': '6', 'color': '999999'})
        set_cell_shading(c_draw, "FFFFFF")

        p_dw = c_draw.paragraphs[0]
        p_dw.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dw.paragraph_format.space_before = Pt(8)
        p_dw.paragraph_format.space_after = Pt(45)
        rdw = p_dw.add_run("(Dibuja aquí el símbolo)")
        rdw.font.name = "Arial"
        rdw.font.size = Pt(8)
        rdw.italic = True
        rdw.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)

    out_path = os.path.join(OUTPUT_DIR, "Evaluacion_Final_Ciencias_5Basico.docx")
    doc.save(out_path)
    print(f"Estudiante test updated: {out_path}")

if __name__ == "__main__":
    build_evaluacion_estudiante()
