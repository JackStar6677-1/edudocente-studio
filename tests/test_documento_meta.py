# -*- coding: utf-8 -*-
"""Todo documento generado sale en español de Chile (es-CL) y con el autor configurado (no "python-docx").

    python -m unittest tests.test_documento_meta -v
"""
import os
import shutil
import tempfile
import unittest
import zipfile

import docx

import documento_meta
from engine import DocenteEngine
from planificacion import anual, mensual, temario


def _xml(ruta, parte):
    with zipfile.ZipFile(ruta) as z:
        return z.read(parte).decode('utf-8')


class PruebasMetadatos(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self._previo = os.environ.get('EDUDOCENTE_AUTOR')
        os.environ['EDUDOCENTE_AUTOR'] = 'Autor de Prueba'

    def tearDown(self):
        if self._previo is None:
            os.environ.pop('EDUDOCENTE_AUTOR', None)
        else:
            os.environ['EDUDOCENTE_AUTOR'] = self._previo
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _verificar(self, ruta):
        core = _xml(ruta, 'docProps/core.xml')
        self.assertIn('<dc:creator>Autor de Prueba</dc:creator>', core)
        self.assertIn('<cp:lastModifiedBy>Autor de Prueba</cp:lastModifiedBy>', core)
        self.assertNotIn('python-docx', core)
        estilos = _xml(ruta, 'word/styles.xml')
        self.assertIn('w:val="es-CL"', estilos)
        self.assertNotIn('en-US', estilos)
        self.assertNotIn('en-US', _xml(ruta, 'word/document.xml'))

    def test_planificaciones_y_temario(self):
        plan = {'anio': 2030, 'profesor': 'D', 'curso': '1°A', 'curso_codigo': '1°A', 'asignatura': 'Música',
                'unidad': {'numero': '1', 'nombre': 'X'}, 'horas': '2',
                'oa': [{'codigo': '1', 'texto': 'T', 'puntos': [], 'indicadores': ['I']}],
                'periodos': [{'nombre': 'Marzo', 'clases': mensual.clases_vacias(1)}]}
        self._verificar(mensual.construir(plan, os.path.join(self.tmp, 'm.docx')))
        self._verificar(anual.construir(anual.plan_en_blanco('1°A', '1°A', 'Música', 2030, 'D'), os.path.join(self.tmp, 'a.docx')))
        self._verificar(temario.construir({'curso': '1°A', 'asignatura': 'Música'}, os.path.join(self.tmp, 't.docx')))

    def test_evaluacion(self):
        cfg = {'colegio': 'X', 'asignatura': 'CIENCIAS', 'curso': '5°', 'titulo': 'E', 'puntaje_total': 4,
               'item2_verdadero_falso': [{'oracion': 'a', 'resp': 'V', 'justificacion': 'b'}], 'item2_pts_cada_una': 1}
        ruta = os.path.join(self.tmp, 'e.docx')
        DocenteEngine(cfg).build_evaluacion_estudiante(ruta)
        self._verificar(ruta)

    def test_normalizar_un_docx_existente_conserva_el_contenido(self):
        d = docx.Document()                                   # python-docx puro: autor "python-docx" e idioma ingles
        d.add_paragraph('Escuchar y apreciar música')
        ruta = os.path.join(self.tmp, 'viejo.docx')
        d.save(ruta)
        self.assertIn('python-docx', _xml(ruta, 'docProps/core.xml'))
        documento_meta.normalizar_docx(ruta)
        self._verificar(ruta)
        self.assertEqual(docx.Document(ruta).paragraphs[0].text, 'Escuchar y apreciar música')

    def test_autor_por_defecto_y_auto(self):
        os.environ.pop('EDUDOCENTE_AUTOR')
        self.assertEqual(documento_meta.autor(), 'Jack')       # valor por defecto del proyecto (config/documento.json)
        os.environ['EDUDOCENTE_AUTOR'] = 'auto'
        self.assertTrue(documento_meta.autor())                 # nombre de usuario del equipo


if __name__ == '__main__':
    unittest.main()
