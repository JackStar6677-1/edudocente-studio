import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

from institution_manager import institution_manager

# Carga dinámica de identidad institucional activa
active_theme = institution_manager.get_active_theme_bundle()
COLOR_NAVY_HEX = active_theme["primary_hex"]
COLOR_ICE_HEX = active_theme["secondary_hex"]
COLOR_BOX_BG_HEX = active_theme["card_bg_hex"]
COLOR_BORDER_HEX = active_theme["border_hex"]
COLOR_CORRECT_HEX = active_theme["teacher_correct_hex"]

COLOR_NAVY_RGB = active_theme["primary_rgb"]
COLOR_WHITE_RGB = active_theme["white_rgb"]
COLOR_DARK_RGB = active_theme["dark_rgb"]
COLOR_GRAY_RGB = active_theme["gray_rgb"]
COLOR_CORRECT_RGB = active_theme["teacher_correct_rgb"]

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = active_theme["logo_path"]
ASSETS_DIR = os.path.join(BASE_DIR, "assets", "circuit_components")
MATTER_ASSETS_DIR = os.path.join(BASE_DIR, "assets", "matter_states")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
DOWNLOADS_DIR = r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita"
os.makedirs(OUTPUT_DIR, exist_ok=True)

def reload_active_theme():
    """Recarga los colores y logo cuando el usuario cambia de institución"""
    global COLOR_NAVY_HEX, COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_CORRECT_HEX
    global COLOR_NAVY_RGB, COLOR_WHITE_RGB, COLOR_DARK_RGB, COLOR_GRAY_RGB, COLOR_CORRECT_RGB, LOGO_PATH
    theme = institution_manager.get_active_theme_bundle()
    COLOR_NAVY_HEX = theme["primary_hex"]
    COLOR_ICE_HEX = theme["secondary_hex"]
    COLOR_BOX_BG_HEX = theme["card_bg_hex"]
    COLOR_BORDER_HEX = theme["border_hex"]
    COLOR_CORRECT_HEX = theme["teacher_correct_hex"]
    COLOR_NAVY_RGB = theme["primary_rgb"]
    COLOR_WHITE_RGB = theme["white_rgb"]
    COLOR_DARK_RGB = theme["dark_rgb"]
    COLOR_GRAY_RGB = theme["gray_rgb"]
    COLOR_CORRECT_RGB = theme["teacher_correct_rgb"]
    LOGO_PATH = theme["logo_path"]
    return theme

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_cell_borders(cell, **kwargs):
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            val = edge_data.get('val', 'single')
            sz = edge_data.get('sz', '4')
            color = edge_data.get('color', COLOR_BORDER_HEX)
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    cell._tc.get_or_add_tcPr().append(tcBorders)

def set_document_language(doc, lang_code="es-CL"):
    """
    Configura el idioma predeterminado de corrección ortográfica y gramatical
    en los metadatos de estilos OpenXML (w:docDefaults y estilo Normal).
    Evita que Microsoft Word marque todo el texto en español con líneas rojas de error.
    """
    try:
        styles_elm = doc.styles.element
        doc_defaults = styles_elm.find(qn('w:docDefaults'))
        if doc_defaults is not None:
            rpr_default = doc_defaults.find(qn('w:rPrDefault'))
            if rpr_default is not None:
                rpr = rpr_default.find(qn('w:rPr'))
                if rpr is not None:
                    lang = rpr.find(qn('w:lang'))
                    if lang is not None:
                        lang.set(qn('w:val'), lang_code)
                        lang.set(qn('w:eastAsia'), lang_code)
                        lang.set(qn('w:bidi'), 'ar-SA')
                    else:
                        lang = OxmlElement('w:lang')
                        lang.set(qn('w:val'), lang_code)
                        lang.set(qn('w:eastAsia'), lang_code)
                        lang.set(qn('w:bidi'), 'ar-SA')
                        rpr.append(lang)

        # Configurar explícitamente en el estilo Normal
        normal = doc.styles['Normal']
        normal_rpr = normal.element.get_or_add_rPr()
        normal_lang = normal_rpr.find(qn('w:lang'))
        if normal_lang is None:
            normal_lang = OxmlElement('w:lang')
            normal_rpr.append(normal_lang)
        normal_lang.set(qn('w:val'), lang_code)
        normal_lang.set(qn('w:eastAsia'), lang_code)
        normal_lang.set(qn('w:bidi'), 'ar-SA')
    except Exception as e:
        print(f"[Aviso] No se pudo configurar metadatos de idioma: {e}")

def create_base_doc(lang_code="es-CL"):
    doc = docx.Document()
    set_document_language(doc, lang_code=lang_code)
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
    return doc

def safe_save(doc, path):
    """Guarda el documento Word capturando posibles bloqueos por tener el archivo abierto en Word."""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    try:
        doc.save(path)
        return True
    except PermissionError:
        print(f"  [AVISO] No se pudo sobrescribir '{os.path.basename(path)}' en '{os.path.dirname(path)}' porque el archivo está abierto en Microsoft Word.")
        return False
    except Exception as e:
        print(f"  [ERROR] Al guardar '{path}': {e}")
        return False

def add_header(doc, school=None, subject="CIENCIAS NATURALES - 5° BÁSICO A - B", subtitle=None):
    reload_active_theme()
    active_inst = institution_manager.get_active()
    effective_school = school if school else active_inst.get("name", "COLEGIO LUIS PASTEUR ANEXO")

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Logo
    cell_logo = table.cell(0, 0)
    cell_logo.width = Pt(58.0)
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(0)
    if os.path.exists(LOGO_PATH):
        r_logo = p_logo.add_run()
        r_logo.add_picture(LOGO_PATH, width=Inches(0.62))

    # School & Subject
    cell_title = table.cell(0, 1)
    cell_title.width = Pt(453.2)
    set_cell_shading(cell_title, COLOR_NAVY_HEX)
    set_cell_margins(cell_title, top=120, bottom=120, left=180, right=140)
    cell_title.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p_title = cell_title.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(0)

    r1 = p_title.add_run(f"{effective_school}\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.bold = True
    r1.font.color.rgb = COLOR_WHITE_RGB

    text_sub = subject if not subtitle else f"{subject}  |  {subtitle}"
    r2 = p_title.add_run(text_sub)
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.bold = True
    r2.font.color.rgb = COLOR_WHITE_RGB

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(2)

def add_title_banner(doc, title_text, is_pauta=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Pt(511.2)
    fill_col = "E2F0D9" if is_pauta else COLOR_ICE_HEX
    text_col = RGBColor(0x1B, 0x5E, 0x20) if is_pauta else COLOR_NAVY_RGB

    set_cell_shading(cell, fill_col)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run(title_text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = text_col

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(2)

def add_student_info(doc, is_pauta=False):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    c00, c01 = table.cell(0, 0), table.cell(0, 1)
    c10, c11 = table.cell(1, 0), table.cell(1, 1)

    c00.width = Pt(340)
    c01.width = Pt(171.2)
    c10.width = Pt(340)
    c11.width = Pt(171.2)

    for c in [c00, c01, c10, c11]:
        set_cell_margins(c, top=70, bottom=70, left=90, right=90)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

    # Fila 0
    p = c00.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if is_pauta:
        r = p.add_run("DOCUMENTO DOCENTE: PAUTA DE CORRECCIÓN")
        r.font.color.rgb = COLOR_CORRECT_RGB
    else:
        r = p.add_run("Nombre: __________________________________________________")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.bold = True

    p = c01.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Curso: 5° Básico [ A ]  [ B ]\nFecha: _____ / _____ / 2026")
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.bold = True

    # Fila 1
    p = c10.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("Objetivo (OA 11): ")
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.bold = True
    r2 = p.add_run("Explicar la importancia de la energía eléctrica en la vida cotidiana y proponer medidas para promover su ahorro y uso responsable.")
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)

    p = c11.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if is_pauta:
        r = p.add_run("Puntaje Total: 28 puntos\nAsignatura: Ciencias Naturales")
    else:
        r = p.add_run("Puntaje Total: 28 puntos\nPuntaje Obtenido: _____   Nota: _____")
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.bold = True

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def add_instructions_box(doc):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    c = t.cell(0, 0)
    c.width = Pt(511.2)
    set_cell_shading(c, "F6FAFE")
    set_cell_margins(c, top=70, bottom=70, left=110, right=110)
    set_cell_borders(c, left={'val': 'single', 'sz': '16', 'color': COLOR_NAVY_HEX})

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r_bold = p.add_run("Instrucciones: ")
    r_bold.font.name = "Arial"
    r_bold.font.size = Pt(8.5)
    r_bold.bold = True
    r_bold.font.color.rgb = COLOR_NAVY_RGB

    r_text = p.add_run("Lee atentamente cada pregunta antes de responder. Marca con una ")
    r_text.font.name = "Arial"
    r_text.font.size = Pt(8.5)

    r_x = p.add_run("X")
    r_x.bold = True
    r_x.font.name = "Arial"
    r_x.font.size = Pt(8.5)

    r_end = p.add_run(" la alternativa que consideres correcta. Revisa muy bien tus respuestas antes de entregar tu evaluación.")
    r_end.font.name = "Arial"
    r_end.font.size = Pt(8.5)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def add_section_header(doc, title):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Pt(511.2)
    set_cell_shading(cell, COLOR_NAVY_HEX)
    set_cell_margins(cell, top=80, bottom=80, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = COLOR_WHITE_RGB

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(3)
