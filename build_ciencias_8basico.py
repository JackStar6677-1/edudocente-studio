"""
Generador de Evaluación Final y Pauta de Corrección: Ciencias Naturales 8° Básico A
Colegio Luis Pasteur Anexo - Profesora Margarita Miranda B.
Año Escolar 2026
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

DOWNLOADS_ROOT = r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita"
DEST_8BASICO = os.path.join(DOWNLOADS_ROOT, "01 - Ciencias Naturales", "8° Básico (8°A)")
os.makedirs(DEST_8BASICO, exist_ok=True)

ATOM_ASSETS_DIR = os.path.join(BASE_DIR, "assets", "atom_model")

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_section_header,
    set_cell_shading, set_cell_borders, set_cell_margins, COLOR_NAVY_HEX,
    COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_CORRECT_HEX,
    COLOR_NAVY_RGB, COLOR_CORRECT_RGB, COLOR_DARK_RGB, COLOR_WHITE_RGB, safe_save
)

def add_student_info_8b(doc, is_pauta=False):
    t_info = doc.add_table(rows=2, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False
    w_left, w_right = Pt(255), Pt(256.2)

    filas = [
        ("Nombre del/la Estudiante: " + ("PAUTA OFICIAL DOCENTE" if is_pauta else "____________________________________"),
         "Curso: 8° Básico A"),
        ("Fecha: ______ de _________________ de 2026",
         "Puntaje Total: 25 puntos   |   Puntaje Obtenido: " + ("25 pts" if is_pauta else "_____ pts"))
    ]

    for r_idx, (d1, d2) in enumerate(filas):
        c1, c2 = t_info.cell(r_idx, 0), t_info.cell(r_idx, 1)
        c1.width, c2.width = w_left, w_right
        for c, text in [(c1, d1), (c2, d2)]:
            set_cell_margins(c, top=55, bottom=55, left=80, right=80)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            if r_idx == 0:
                set_cell_shading(c, "EAF9E6" if is_pauta else COLOR_ICE_HEX)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if is_pauta and "PAUTA" in text:
                r.bold = True
                r.font.color.rgb = COLOR_CORRECT_RGB
            elif "Puntaje" in text or "Curso" in text:
                r.bold = True

    # Cuadro de instrucciones pedagógicas
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    t_inst = doc.add_table(rows=1, cols=1)
    t_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inst.autofit = False
    c_ins = t_inst.cell(0, 0)
    c_ins.width = Pt(511.2)
    set_cell_shading(c_ins, COLOR_BOX_BG_HEX)
    set_cell_margins(c_ins, top=50, bottom=50, left=80, right=80)
    set_cell_borders(c_ins, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'val': 'single', 'sz': '12', 'color': COLOR_NAVY_HEX},
                            right={'color': COLOR_BORDER_HEX})
    p_ins = c_ins.paragraphs[0]
    p_ins.paragraph_format.space_before = Pt(0)
    p_ins.paragraph_format.space_after = Pt(0)

    r_oa = p_ins.add_run("Objetivo de Aprendizaje (OA 3): ")
    r_oa.bold = True
    r_oa.font.name = "Arial"
    r_oa.font.size = Pt(8.5)
    r_oa.font.color.rgb = COLOR_NAVY_RGB

    r_oat = p_ins.add_run("Desarrollar modelos que expliquen que la materia está constituida por átomos que interactúan generando diversas partículas y sustancias.\n")
    r_oat.font.name = "Arial"
    r_oat.font.size = Pt(8.5)

    r_cont = p_ins.add_run("Contenidos Evaluados: ")
    r_cont.bold = True
    r_cont.font.name = "Arial"
    r_cont.font.size = Pt(8)
    r_contt = p_ins.add_run("Teoría atómica de Dalton (Pág. 126), estructura del átomo y partículas subatómicas (Págs. 136-137).\n")
    r_contt.font.name = "Arial"
    r_contt.font.size = Pt(8)

    r_txt = p_ins.add_run(
        "Instrucciones: Lee atentamente cada ítem antes de responder. Marca con claridad tus respuestas "
        "y utiliza letra legible. Si tienes dudas, levanta la mano para consultar con la profesora."
    )
    r_txt.font.name = "Arial"
    r_txt.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ==============================================================================
# EVALUACIÓN ESTUDIANTE (8° BÁSICO)
# ==============================================================================
def build_evaluacion_estudiante_8b():
    doc = create_base_doc()

    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="CIENCIAS NATURALES — 8° BÁSICO A")
    add_title_banner(doc, "EVALUACIÓN FINAL: EL ÁTOMO Y LA TEORÍA ATÓMICA DE DALTON")
    add_student_info_8b(doc, is_pauta=False)

    # ------------------ ÍTEM I: SELECCIÓN MÚLTIPLE ------------------
    add_section_header(doc, "ÍTEM I: SELECCIÓN MÚLTIPLE (2 puntos c/u — Total: 10 puntos)")

    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(1)
    p_inst.paragraph_format.space_after = Pt(4)
    r_inst = p_inst.add_run("Lee atentamente cada pregunta y marca con una ")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(8.5)
    rx = p_inst.add_run("X")
    rx.bold = True
    rx.font.name = "Arial"
    rx.font.size = Pt(8.5)
    r_inst2 = p_inst.add_run(" la alternativa que consideres correcta:")
    r_inst2.font.name = "Arial"
    r_inst2.font.size = Pt(8.5)

    preguntas_item1 = [
        ("1. ¿Cuáles fueron los principales postulados formulados por John Dalton (1766-1844)?",
         [("A", "La materia está compuesta por diminutas partículas indivisibles llamadas átomos."),
          ("B", "La materia está formada exclusivamente por cargas magnéticas e imanes."),
          ("C", "La materia es únicamente todo aquello que posee energía lumínica y espacial."),
          ("D", "La materia está compuesta solamente por los seres vivos y organismos orgánicos.")]),

        ("2. Según la teoría atómica de Dalton, la materia está formada por diminutas partículas:",
         [("A", "Visibles a simple vista y destructibles llamadas átomos."),
          ("B", "Divisibles en fragmentos más pequeños y destructibles llamadas átomos."),
          ("C", "Indivisibles e indestructibles llamadas átomos."),
          ("D", "Estructuradas en capas planas y alterables llamadas átomos.")]),

        ("3. ¿Cómo concebía el átomo John Dalton en su modelo científico?",
         [("A", "Como un círculo geométrico plano sin masa."),
          ("B", "Como una esfera maciza, compacta e indivisible."),
          ("C", "Como un triángulo plano que absorbe radiación."),
          ("D", "Como un rectángulo hueco de materia fluida.")]),

        ("4. Uno de los aportes fundamentales de Dalton a la teoría atómica moderna establece que:",
         [("A", "La materia está compuesta por átomos que se reorganizan sin perder masa en las reacciones químicas."),
          ("B", "Los átomos se combinan siguiendo estrictamente la sucesión de números primos."),
          ("C", "En cada reacción química siempre ocurre una destrucción total de la masa atómica."),
          ("D", "La materia está formada por círculos planos concéntricos que cambian de color.")]),

        ("5. Con los avances de la ciencia, hoy sabemos que los átomos están formados por:",
         [("A", "Partículas de protones, neutrón y núcleo sin masa."),
          ("B", "Tres partículas subatómicas fundamentales: protones, neutrones y electrones."),
          ("C", "Tres esferas macizas clasificadas en mediana, grande y pequeña."),
          ("D", "Tres elementos primarios de la naturaleza: viento, agua y radiación solar.")])
    ]

    for num, alternativas in preguntas_item1:
        t_q = doc.add_table(rows=1, cols=1)
        t_q.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_q.autofit = False
        c_q = t_q.cell(0, 0)
        c_q.width = Pt(511.2)
        set_cell_shading(c_q, COLOR_BOX_BG_HEX)
        set_cell_margins(c_q, top=45, bottom=45, left=60, right=60)
        set_cell_borders(c_q, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                               left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
        pq = c_q.paragraphs[0]
        pq.paragraph_format.space_before = Pt(0)
        pq.paragraph_format.space_after = Pt(0)
        rq = pq.add_run(num)
        rq.bold = True
        rq.font.name = "Arial"
        rq.font.size = Pt(8.5)
        rq.font.color.rgb = COLOR_NAVY_RGB

        t_opts = doc.add_table(rows=4, cols=1)
        t_opts.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_opts.autofit = False

        for idx_opt, (letra, texto_alt) in enumerate(alternativas):
            cell_opt = t_opts.cell(idx_opt, 0)
            cell_opt.width = Pt(511.2)
            set_cell_margins(cell_opt, top=30, bottom=30, left=50, right=50)
            set_cell_borders(cell_opt, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})
            p_opt = cell_opt.paragraphs[0]
            p_opt.paragraph_format.space_before = Pt(0)
            p_opt.paragraph_format.space_after = Pt(0)

            r_letra = p_opt.add_run(f"{letra})  ")
            r_letra.bold = True
            r_letra.font.name = "Arial"
            r_letra.font.size = Pt(8.5)
            r_letra.font.color.rgb = COLOR_NAVY_RGB

            r_alt = p_opt.add_run(texto_alt)
            r_alt.font.name = "Arial"
            r_alt.font.size = Pt(8.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------ ÍTEM II: VERDADERO O FALSO ------------------
    add_section_header(doc, "ÍTEM II: VERDADERO O FALSO (1 punto c/u — Total: 10 puntos)")

    p_vf_inst = doc.add_paragraph()
    p_vf_inst.paragraph_format.space_before = Pt(1)
    p_vf_inst.paragraph_format.space_after = Pt(4)
    r_vfi = p_vf_inst.add_run("De las siguientes oraciones, escribe en el paréntesis una ")
    r_vfi.font.name = "Arial"
    r_vfi.font.size = Pt(8.5)
    r_v = p_vf_inst.add_run("V")
    r_v.bold = True
    r_v.font.name = "Arial"
    r_v.font.size = Pt(8.5)
    r_vfi2 = p_vf_inst.add_run(" si la afirmación es verdadera o una ")
    r_vfi2.font.name = "Arial"
    r_vfi2.font.size = Pt(8.5)
    r_f = p_vf_inst.add_run("F")
    r_f.bold = True
    r_f.font.name = "Arial"
    r_f.font.size = Pt(8.5)
    r_vfi3 = p_vf_inst.add_run(" si es falsa:")
    r_vfi3.font.name = "Arial"
    r_vfi3.font.size = Pt(8.5)

    afirmaciones_vf = [
        "El átomo está formado por electrones, protones y neutrones.",
        "La notación atómica se utiliza para representar la composición del átomo.",
        "La notación atómica incluye el número atómico (Z) y el número másico (A).",
        "El número atómico corresponde a la cantidad de esferas del átomo.",
        "El número másico representa el número total de protones y neutrones presentes en el núcleo.",
        "El número másico se representa matemáticamente como A = P + N (protones + neutrones).",
        "Los átomos son rectángulos indestructibles que forman la materia.",
        "El átomo está formado únicamente por neutrones.",
        "Los átomos forman círculos destructibles durante los cambios físicos.",
        "El número total de protones en todos los átomos siempre es exactamente 10."
    ]

    t_vf = doc.add_table(rows=len(afirmaciones_vf), cols=3)
    t_vf.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_vf.autofit = False

    for idx, texto in enumerate(afirmaciones_vf, 1):
        c_num = t_vf.cell(idx - 1, 0)
        c_box = t_vf.cell(idx - 1, 1)
        c_txt = t_vf.cell(idx - 1, 2)

        c_num.width = Pt(22)
        c_box.width = Pt(40)
        c_txt.width = Pt(449.2)

        bg_col = "FFFFFF" if idx % 2 != 0 else COLOR_BOX_BG_HEX

        for c in (c_num, c_box, c_txt):
            set_cell_shading(c, bg_col)
            set_cell_margins(c, top=35, bottom=35, left=35, right=35)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

        pn = c_num.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(0)
        pn.paragraph_format.space_after = Pt(0)
        rn = pn.add_run(f"{idx}.")
        rn.font.name = "Arial"
        rn.font.size = Pt(8.5)
        rn.bold = True
        rn.font.color.rgb = COLOR_NAVY_RGB

        pb = c_box.paragraphs[0]
        pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pb.paragraph_format.space_before = Pt(0)
        pb.paragraph_format.space_after = Pt(0)
        rb = pb.add_run("(     )")
        rb.font.name = "Arial"
        rb.font.size = Pt(8.5)
        rb.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        pt = c_txt.paragraphs[0]
        pt.paragraph_format.space_before = Pt(0)
        pt.paragraph_format.space_after = Pt(0)
        rt = pt.add_run(texto)
        rt.font.name = "Arial"
        rt.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # ------------------ ÍTEM III: DIBUJO Y ROTULACIÓN DEL ÁTOMO ------------------
    add_section_header(doc, "ÍTEM III: DIBUJO, ROTULACIÓN Y PINTADO DEL ÁTOMO (5 puntos)")

    p_i3 = doc.add_paragraph()
    p_i3.paragraph_format.space_before = Pt(1)
    p_i3.paragraph_format.space_after = Pt(4)
    r_i3 = p_i3.add_run(
        "Observa el siguiente esquema de un átomo. Completa las líneas escribiendo el nombre de las partes "
        "indicadas por las flechas (Núcleo, Corteza, Protón, Neutrón y Electrón). Luego, dibuja las partículas "
        "faltantes y pinta según lo visto en clases:"
    )
    r_i3.font.name = "Arial"
    r_i3.font.size = Pt(8.5)

    # Insertar imagen del esqueleto
    skel_path = os.path.join(ATOM_ASSETS_DIR, "atom_skeleton_student.png")
    if os.path.exists(skel_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(skel_path, width=Inches(5.6))

    # Cuadro de criterios de logro
    t_crit = doc.add_table(rows=1, cols=1)
    t_crit.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_crit.autofit = False
    c_cr = t_crit.cell(0, 0)
    c_cr.width = Pt(511.2)
    set_cell_shading(c_cr, COLOR_BOX_BG_HEX)
    set_cell_margins(c_cr, top=40, bottom=40, left=70, right=70)
    set_cell_borders(c_cr, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                           left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
    p_cr = c_cr.paragraphs[0]
    p_cr.paragraph_format.space_before = Pt(0)
    p_cr.paragraph_format.space_after = Pt(0)
    rc1 = p_cr.add_run("Criterios de evaluación (1 punto c/u — Total 5 pts): ")
    rc1.bold = True
    rc1.font.name = "Arial"
    rc1.font.size = Pt(8)
    rc1.font.color.rgb = COLOR_NAVY_RGB
    rc2 = p_cr.add_run(
        "1. Identifica y rotula la Corteza/Órbita. | 2. Identifica y rotula el Electrón (e–). | "
        "3. Identifica y rotula el Núcleo atómico. | 4. Identifica y rotula el Protón (p+). | "
        "5. Identifica y rotula el Neutrón (n0) y colorea adecuadamente."
    )
    rc2.font.name = "Arial"
    rc2.font.size = Pt(7.5)

    out_name = "Evaluacion_Final_Ciencias_8Basico.docx"
    safe_save(doc, os.path.join(OUTPUT_DIR, out_name))
    safe_save(doc, os.path.join(DEST_8BASICO, out_name))
    print(f"  [OK] Prueba Estudiante 8° Básico generada: {out_name}")

# ==============================================================================
# PAUTA DE CORRECCIÓN (DOCENTE 8° BÁSICO)
# ==============================================================================
def build_pauta_correccion_8b():
    doc = create_base_doc()

    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="CIENCIAS NATURALES — 8° BÁSICO A", subtitle="DOCUMENTO DOCENTE")
    add_title_banner(doc, "PAUTA DE CORRECCIÓN: EVALUACIÓN FINAL DE CIENCIAS NATURALES", is_pauta=True)
    add_student_info_8b(doc, is_pauta=True)

    # ------------------ SOLUCIONARIO ÍTEM I ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM I: SELECCIÓN MÚLTIPLE (2 pts c/u — Total: 10 pts)")

    sol_item1 = [
        ("1. ¿Cuáles fueron los principales postulados formulados por John Dalton (1766-1844)?",
         "A) La materia está compuesta por diminutas partículas indivisibles llamadas átomos.",
         "Justificación: Dalton propuso el primer modelo atómico con bases científicas (1808), postulando que toda la materia está constituida por partículas fundamentales indivisibles e indestructibles (Texto escolar Pág. 126)."),

        ("2. Según la teoría atómica de Dalton, la materia está formada por diminutas partículas:",
         "C) Indivisibles e indestructibles llamadas átomos.",
         "Justificación: Para Dalton el átomo representaba la unidad última de la materia, postulando que los átomos no podían dividirse ni destruirse, conservándose inalterables en las reacciones (Texto escolar Pág. 126)."),

        ("3. ¿Cómo concebía el átomo John Dalton en su modelo científico?",
         "B) Como una esfera maciza, compacta e indivisible.",
         "Justificación: El modelo corpuscular de Dalton visualizaba los átomos como diminutas bolas de billar o esferas macizas, homogéneas y sin estructura interna descubierta aún en su época (Texto escolar Pág. 126)."),

        ("4. Uno de los aportes fundamentales de Dalton a la teoría atómica moderna establece que:",
         "A) La materia está compuesta por átomos que se reorganizan sin perder masa en las reacciones químicas.",
         "Justificación: Dalton explicó la Ley de Conservación de la Masa: en una reacción química los átomos no se crean ni se destruyen, únicamente se redistribuyen formando nuevos compuestos (Texto escolar Pág. 126)."),

        ("5. Con los avances de la ciencia, hoy sabemos que los átomos están formados por:",
         "B) Tres partículas subatómicas fundamentales: protones, neutrones y electrones.",
         "Justificación: El modelo científico actual demuestra que el átomo es divisible y se compone de un núcleo con protones y neutrones, y una electrosfera o corteza con electrones (Texto escolar Págs. 136-137).")
    ]

    for preg, resp, just in sol_item1:
        t_s = doc.add_table(rows=1, cols=1)
        t_s.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_s.autofit = False
        cs = t_s.cell(0, 0)
        cs.width = Pt(511.2)
        set_cell_shading(cs, "F6FAFE")
        set_cell_margins(cs, top=45, bottom=45, left=60, right=60)
        set_cell_borders(cs, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                               left={'val': 'single', 'sz': '12', 'color': COLOR_CORRECT_HEX},
                               right={'color': COLOR_BORDER_HEX})
        ps = cs.paragraphs[0]
        ps.paragraph_format.space_before = Pt(0)
        ps.paragraph_format.space_after = Pt(0)

        rp = ps.add_run(f"{preg}\n")
        rp.bold = True
        rp.font.name = "Arial"
        rp.font.size = Pt(8.5)
        rp.font.color.rgb = COLOR_NAVY_RGB

        rr = ps.add_run(f"Respuesta Correcta: {resp}\n")
        rr.bold = True
        rr.font.name = "Arial"
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = COLOR_CORRECT_RGB

        rj = ps.add_run(f"{just}")
        rj.font.name = "Arial"
        rj.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------ SOLUCIONARIO ÍTEM II ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM II: VERDADERO O FALSO (1 pto c/u — Total: 10 pts)")

    sol_vf = [
        ("V", "El átomo está formado por electrones, protones y neutrones.",
         "Es verdadera. Son las tres partículas subatómicas fundamentales."),

        ("V", "La notación atómica se utiliza para representar la composición del átomo.",
         "Es verdadera. Informa el símbolo químico, número atómico y número másico."),

        ("V", "La notación atómica incluye el número atómico (Z) y el número másico (A).",
         "Es verdadera. Se escribe con A como superíndice y Z como subíndice izquierdo."),

        ("F", "El número atómico corresponde a la cantidad de esferas del átomo.",
         "Es falsa. El número atómico (Z) corresponde exactamente a la cantidad de protones."),

        ("V", "El número másico representa el número total de protones y neutrones presentes en el núcleo.",
         "Es verdadera. Define la masa aproximada del átomo concentrada en su núcleo."),

        ("V", "El número másico se representa matemáticamente como A = P + N (protones + neutrones).",
         "Es verdadera. A = Z + n0."),

        ("F", "Los átomos son rectángulos indestructibles que forman la materia.",
         "Es falsa. Son entidades con núcleo y nube electrónica, no figuras geométricas rectangulares."),

        ("F", "El átomo está formado únicamente por neutrones.",
         "Es falsa. Contiene también protones (positivos) y electrones (negativos)."),

        ("F", "Los átomos forman círculos destructibles durante los cambios físicos.",
         "Es falsa. En los cambios físicos los átomos conservan intacta su estructura nuclear."),

        ("F", "El número total de protones en todos los átomos siempre es exactamente 10.",
         "Es falsa. Cada elemento posee un número atómico propio (ej. Hidrógeno tiene 1, Carbono 6, etc.).")
    ]

    t_vf_s = doc.add_table(rows=len(sol_vf) + 1, cols=4)
    t_vf_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_vf_s.autofit = False

    w_cols = [Pt(25), Pt(45), Pt(230), Pt(211.2)]
    headers_vf = ["N°", "Resp.", "Afirmación Evaluada", "Justificación Pedagógica"]

    for i, h in enumerate(headers_vf):
        c = t_vf_s.cell(0, i)
        c.width = w_cols[i]
        set_cell_shading(c, COLOR_NAVY_HEX)
        set_cell_margins(c, top=45, bottom=45, left=45, right=45)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = COLOR_WHITE_RGB

    for row_idx, (letra, afirmacion, justificacion) in enumerate(sol_vf, 1):
        bg = "FFFFFF" if row_idx % 2 != 0 else COLOR_BOX_BG_HEX
        c_n = t_vf_s.cell(row_idx, 0)
        c_l = t_vf_s.cell(row_idx, 1)
        c_a = t_vf_s.cell(row_idx, 2)
        c_j = t_vf_s.cell(row_idx, 3)

        for col_idx, cell in enumerate((c_n, c_l, c_a, c_j)):
            cell.width = w_cols[col_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=35, bottom=35, left=40, right=40)
            set_cell_borders(cell, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                   left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

        pn = c_n.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(0)
        pn.paragraph_format.space_after = Pt(0)
        rn = pn.add_run(f"{row_idx}")
        rn.font.name = "Arial"
        rn.font.size = Pt(8)
        rn.bold = True

        pl = c_l.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pl.paragraph_format.space_before = Pt(0)
        pl.paragraph_format.space_after = Pt(0)
        rl = pl.add_run(f"[ {letra} ]")
        rl.font.name = "Arial"
        rl.font.size = Pt(8.5)
        rl.bold = True
        rl.font.color.rgb = COLOR_CORRECT_RGB

        pa = c_a.paragraphs[0]
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after = Pt(0)
        ra = pa.add_run(afirmacion)
        ra.font.name = "Arial"
        ra.font.size = Pt(8)

        pj = c_j.paragraphs[0]
        pj.paragraph_format.space_before = Pt(0)
        pj.paragraph_format.space_after = Pt(0)
        rj = pj.add_run(justificacion)
        rj.font.name = "Arial"
        rj.font.size = Pt(7.5)
        rj.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # ------------------ SOLUCIONARIO ÍTEM III ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM III: MODELO ATÓMICO RESUELTO Y RÚBRICA (5 pts)")

    p_sol_img = doc.add_paragraph()
    p_sol_img.paragraph_format.space_before = Pt(1)
    p_sol_img.paragraph_format.space_after = Pt(4)
    r_sim = p_sol_img.add_run("Representación gráfica esperada y rotulación normalizada:")
    r_sim.bold = True
    r_sim.font.name = "Arial"
    r_sim.font.size = Pt(8.5)
    r_sim.font.color.rgb = COLOR_CORRECT_RGB

    solved_path = os.path.join(ATOM_ASSETS_DIR, "atom_solved_teacher.png")
    if os.path.exists(solved_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(solved_path, width=Inches(5.6))

    # Pauta de corrección paso a paso
    pautas_criterios = [
        ("1. Corteza o Electrosfera (1 pto)", "Rotula correctamente la zona exterior por donde orbitan los electrones."),
        ("2. Electrones (e–) (1 pto)", "Identifica y dibuja/pinta las partículas de carga negativa situadas en las órbitas."),
        ("3. Núcleo Atómico (1 pto)", "Rotula con precisión la región central compacta y de carga positiva del átomo."),
        ("4. Protones (p+) (1 pto)", "Identifica y dibuja/pinta las partículas de carga positiva ubicadas dentro del núcleo."),
        ("5. Neutrones (n0) y Pintado (1 pto)", "Identifica partículas neutras en el núcleo y aplica colores diferenciados según lo visto en clases.")
    ]

    t_paut = doc.add_table(rows=len(pautas_criterios), cols=1)
    t_paut.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_paut.autofit = False

    for idx, (crit_t, crit_d) in enumerate(pautas_criterios):
        c = t_paut.cell(idx, 0)
        c.width = Pt(511.2)
        set_cell_shading(c, "F6FAFE" if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c, top=35, bottom=35, left=60, right=60)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'val': 'single', 'sz': '12', 'color': COLOR_CORRECT_HEX},
                            right={'color': COLOR_BORDER_HEX})
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        rct = p.add_run(f"• {crit_t}: ")
        rct.bold = True
        rct.font.name = "Arial"
        rct.font.size = Pt(8)
        rct.font.color.rgb = COLOR_CORRECT_RGB
        rcd = p.add_run(crit_d)
        rcd.font.name = "Arial"
        rcd.font.size = Pt(8)

    out_name = "Pauta_Correccion_Ciencias_8Basico.docx"
    safe_save(doc, os.path.join(OUTPUT_DIR, out_name))
    safe_save(doc, os.path.join(DEST_8BASICO, out_name))
    print(f"  [OK] Pauta Docente 8° Básico generada: {out_name}")

def generar_ciencias_8basico_completa():
    print("\n=======================================================")
    print(" GENERANDO EVALUACIÓN Y PAUTA CIENCIAS NATURALES 8° BÁSICO")
    print("=======================================================")
    build_evaluacion_estudiante_8b()
    build_pauta_correccion_8b()
    print(f"\n[COMPLETADO] Material de Ciencias 8° Básico desplegado en: {DEST_8BASICO}\n")

if __name__ == "__main__":
    generar_ciencias_8basico_completa()
