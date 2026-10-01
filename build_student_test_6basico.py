import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from generator_core import (
    create_base_doc, add_header, add_title_banner,
    add_instructions_box, add_section_header, set_cell_shading,
    set_cell_borders, set_cell_margins, COLOR_NAVY_HEX, COLOR_ICE_HEX,
    COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_DARK_RGB, COLOR_NAVY_RGB,
    OUTPUT_DIR
)

def add_student_info_6b(doc, is_pauta=False):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    c00, c01 = table.cell(0, 0), table.cell(0, 1)
    c10, c11 = table.cell(1, 0), table.cell(1, 1)

    c00.width = Pt(340)
    c01.width = Pt(171.2)
    c10.width = Pt(340)
    c11.width = Pt(171.2)

    for c in [c00, c01, c10, c11]:
        set_cell_margins(c, top=70, bottom=70, left=90, right=90)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

    # Fila 0
    p = c00.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if is_pauta:
        r = p.add_run("DOCUMENTO DOCENTE: PAUTA DE CORRECCIÓN")
        r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
    else:
        r = p.add_run("Nombre: __________________________________________________")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.bold = True

    p = c01.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Curso: 6° Básico [ A ]  [ B ]\nFecha: _____ / _____ / 2026")
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.bold = True

    # Fila 1
    p = c10.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("Objetivo (OA 13): ")
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.bold = True
    r2 = p.add_run("Demostrar, mediante la investigación, los cambios de estado de la materia (fusión, evaporación, ebullición, condensación, solidificación y sublimación).")
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)

    p = c11.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if is_pauta:
        r = p.add_run("Puntaje Total: 26 puntos\nAsignatura: Ciencias Naturales")
    else:
        r = p.add_run("Puntaje Total: 26 puntos\nPuntaje Obtenido: _____   Nota: _____")
    r.font.name = "Arial"
    r.font.size = Pt(8.5)
    r.bold = True

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def build_evaluacion_estudiante_6b():
    doc = create_base_doc()

    # ================== PÁGINA 1 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 6° BÁSICO A - B")
    add_title_banner(doc, "EVALUACIÓN FINAL: CAMBIOS DE ESTADO DE LA MATERIA")
    add_student_info_6b(doc, is_pauta=False)
    add_instructions_box(doc)

    # ÍTEM I
    add_section_header(doc, "ÍTEM I: SELECCIÓN MÚLTIPLE (2 puntos c/u — Total: 10 puntos)")
    
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(1)
    p_inst.paragraph_format.space_after = Pt(4)
    r_inst = p_inst.add_run("Lee atentamente cada pregunta y marca con una ")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(8.5)
    r_x = p_inst.add_run("X")
    r_x.bold = True
    r_x.font.name = "Arial"
    r_x.font.size = Pt(8.5)
    r_inst2 = p_inst.add_run(" la alternativa que consideres correcta.")
    r_inst2.font.name = "Arial"
    r_inst2.font.size = Pt(8.5)

    preguntas_item1 = [
        ("1. Los cambios de estado de la materia son:",
         [("A", "Son los cambios psicológicos de un ser humano."),
          ("B", "Son los cambios físicos que se producen en la pubertad."),
          ("C", "Los cambios físicos en los que la materia cambia su aspecto externo, sin alterar su composición."),
          ("D", "Son cambios en la corteza y el núcleo de la Tierra.")]),

        ("2. En un cambio de estado, la materia solo cambia su aspecto...",
         [("A", "pero continúa siendo la misma sustancia."),
          ("B", "pero no se puede volver a cambiar de estado."),
          ("C", "pero deja de ser una sustancia sólida."),
          ("D", "pero destruye las partículas que la componen.")]),

        ("3. Los cambios de estado se producen por:",
         [("A", "Movimientos de rotación y traslación."),
          ("B", "Por absorción o por liberación de energía (calor)."),
          ("C", "Por técnicas de movimiento moderado."),
          ("D", "Por energía y movimiento de los seres vivos.")]),

        ("4. Los cambios progresivos de la materia son:",
         [("A", "Los cambios por absorción de energía (calor)."),
          ("B", "Los cambios por la corteza terrestre."),
          ("C", "Los cambios por liberación de energía."),
          ("D", "Los cambios energéticos de recursos naturales.")]),

        ("5. Los cambios regresivos de la materia son aquellos que ocurren:",
         [("A", "Por liberación de calor (enfriamiento)."),
          ("B", "Por liberación de agua."),
          ("C", "Por liberación de recursos."),
          ("D", "Por absorción de radiación solar.")])
    ]

    for q_text, options in preguntas_item1:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(3)
        p_q.paragraph_format.space_after = Pt(2)
        rq = p_q.add_run(q_text)
        rq.font.name = "Arial"
        rq.font.size = Pt(9)
        rq.bold = True
        rq.font.color.rgb = COLOR_DARK_RGB

        for letter, opt_text in options:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.2)
            p_opt.paragraph_format.space_before = Pt(0.5)
            p_opt.paragraph_format.space_after = Pt(1)
            
            rl = p_opt.add_run(f"{letter})  ")
            rl.font.name = "Arial"
            rl.font.size = Pt(8.5)
            rl.bold = True
            rl.font.color.rgb = COLOR_NAVY_RGB

            ro = p_opt.add_run(opt_text)
            ro.font.name = "Arial"
            ro.font.size = Pt(8.5)
            ro.font.color.rgb = COLOR_DARK_RGB

    # SALTO A PÁGINA 2
    doc.add_page_break()

    # ================== PÁGINA 2 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 6° BÁSICO A - B", subtitle="EVALUACIÓN FINAL (PÁG. 2)")

    # ÍTEM II
    add_section_header(doc, "ÍTEM II: VERDADERO O FALSO (1 punto c/u — Total: 10 puntos)")

    p_vf_inst = doc.add_paragraph()
    p_vf_inst.paragraph_format.space_before = Pt(1)
    p_vf_inst.paragraph_format.space_after = Pt(4)
    r_vf = p_vf_inst.add_run("Lee atentamente las siguientes oraciones. Escribe una ")
    r_vf.font.name = "Arial"
    r_vf.font.size = Pt(8.5)
    r_v = p_vf_inst.add_run("V")
    r_v.bold = True
    r_v.font.name = "Arial"
    r_v.font.size = Pt(8.5)
    r_vf2 = p_vf_inst.add_run(" si es verdadera o una ")
    r_vf2.font.name = "Arial"
    r_vf2.font.size = Pt(8.5)
    r_f = p_vf_inst.add_run("F")
    r_f.bold = True
    r_f.font.name = "Arial"
    r_f.font.size = Pt(8.5)
    r_vf3 = p_vf_inst.add_run(" si es falsa.")
    r_vf3.font.name = "Arial"
    r_vf3.font.size = Pt(8.5)

    afirmaciones_6b = [
        "1.  ( _____ )  Los cambios de estado por absorción de calor son fusión, vaporización y sublimación.",
        "2.  ( _____ )  El proceso en el cual un líquido pasa a estado gaseoso es la vaporización.",
        "3.  ( _____ )  La fusión es el proceso de un cuerpo en estado sólido a gaseoso.",
        "4.  ( _____ )  El proceso de vaporización puede ocurrir de dos formas: evaporación y ebullición.",
        "5.  ( _____ )  En la ebullición participan pocas o algunas partículas de la superficie del líquido.",
        "6.  ( _____ )  Algunos ejemplos de sustancias que experimentan sublimación son el yodo y la naftalina.",
        "7.  ( _____ )  Los tres estados de la materia estudiados son sólido, líquido y gaseoso.",
        "8.  ( _____ )  Los únicos cambios de estado por liberación de calor son la solidificación y la condensación.",
        "9.  ( _____ )  La solidificación es el proceso en que un líquido pasa a estado sólido.",
        "10. ( _____ )  La condensación es el proceso de un gas cuando pasa a estado líquido."
    ]

    for af in afirmaciones_6b:
        p_af = doc.add_paragraph()
        p_af.paragraph_format.left_indent = Inches(0.15)
        p_af.paragraph_format.space_before = Pt(1.5)
        p_af.paragraph_format.space_after = Pt(2.5)
        r = p_af.add_run(af)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = COLOR_DARK_RGB

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(2)

    # ÍTEM III
    add_section_header(doc, "ÍTEM III: CAMBIOS DE ESTADO POR LIBERACIÓN DE CALOR (Total: 6 puntos)")

    p_i3_inst = doc.add_paragraph()
    p_i3_inst.paragraph_format.space_before = Pt(1)
    p_i3_inst.paragraph_format.space_after = Pt(4)
    r_i3 = p_i3_inst.add_run("Completa con dibujos el estado que se produce por ")
    r_i3.font.name = "Arial"
    r_i3.font.size = Pt(8.5)
    r_i3_b = p_i3_inst.add_run("liberación de calor")
    r_i3_b.bold = True
    r_i3_b.font.name = "Arial"
    r_i3_b.font.size = Pt(8.5)
    r_i3_c = p_i3_inst.add_run(". Dibuja en cada cuadro el estado inicial, la flecha de transformación y el estado final resultante:")
    r_i3_c.font.name = "Arial"
    r_i3_c.font.size = Pt(8.5)

    # 3 bloques consecutivos verticales (Solidificación, Sublimación Inversa, Condensación)
    procesos = [
        {"num": "1", "nombre": "Solidificación", "inicial": "Estado Inicial (Líquido)", "final": "Estado Final (Sólido)", "desc": "Pasa de líquido a sólido por liberación de calor."},
        {"num": "2", "nombre": "Sublimación Inversa (o Regresiva)", "inicial": "Estado Inicial (Gas / Vapor)", "final": "Estado Final (Sólido)", "desc": "Pasa de gas a sólido directamente por liberación de calor."},
        {"num": "3", "nombre": "Condensación", "inicial": "Estado Inicial (Gas / Vapor)", "final": "Estado Final (Líquido)", "desc": "Pasa de gas a líquido por liberación de calor."}
    ]

    for proc in procesos:
        # Fila contenedora
        t_proc = doc.add_table(rows=1, cols=3)
        t_proc.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_proc.autofit = False

        c_ini = t_proc.cell(0, 0)
        c_arrow = t_proc.cell(0, 1)
        c_fin = t_proc.cell(0, 2)

        c_ini.width = Pt(210)
        c_arrow.width = Pt(91.2)
        c_fin.width = Pt(210)

        for c in [c_ini, c_fin]:
            set_cell_margins(c, top=60, bottom=60, left=70, right=70)
            set_cell_borders(c, top={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                bottom={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                left={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                right={'val': 'dashed', 'sz': '6', 'color': '888888'})
            set_cell_shading(c, "FAFCFF")

        set_cell_margins(c_arrow, top=60, bottom=60, left=30, right=30)
        set_cell_borders(c_arrow, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})

        # Cuadro Inicial
        p_in = c_ini.paragraphs[0]
        p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_in.paragraph_format.space_before = Pt(0)
        p_in.paragraph_format.space_after = Pt(2)
        r_in = p_in.add_run(f"Proceso: {proc['nombre']}\n")
        r_in.bold = True
        r_in.font.name = "Arial"
        r_in.font.size = Pt(8.5)
        r_in.font.color.rgb = COLOR_NAVY_RGB
        r_in_sub = p_in.add_run(f"{proc['inicial']}\n\n(Dibuja aquí)\n\n")
        r_in_sub.font.name = "Arial"
        r_in_sub.font.size = Pt(7.5)
        r_in_sub.italic = True
        r_in_sub.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        # Flecha central con texto
        p_ar = c_arrow.paragraphs[0]
        p_ar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ar.paragraph_format.space_before = Pt(8)
        p_ar.paragraph_format.space_after = Pt(0)
        r_ar_txt = p_ar.add_run("Liberación\nde Calor\n\n")
        r_ar_txt.bold = True
        r_ar_txt.font.name = "Arial"
        r_ar_txt.font.size = Pt(7.5)
        r_ar_txt.font.color.rgb = RGBColor(0x17, 0x3F, 0x73)

        r_arrow = p_ar.add_run("────►")
        r_arrow.bold = True
        r_arrow.font.name = "Arial"
        r_arrow.font.size = Pt(12)
        r_arrow.font.color.rgb = RGBColor(0x17, 0x3F, 0x73)

        # Cuadro Final
        p_fn = c_fin.paragraphs[0]
        p_fn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fn.paragraph_format.space_before = Pt(0)
        p_fn.paragraph_format.space_after = Pt(2)
        r_fn = p_fn.add_run(f"Resultado:\n")
        r_fn.bold = True
        r_fn.font.name = "Arial"
        r_fn.font.size = Pt(8.5)
        r_fn.font.color.rgb = COLOR_NAVY_RGB
        r_fn_sub = p_fn.add_run(f"{proc['final']}\n\n(Dibuja aquí)\n\n")
        r_fn_sub.font.name = "Arial"
        r_fn_sub.font.size = Pt(7.5)
        r_fn_sub.italic = True
        r_fn_sub.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        # Separador entre procesos
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(0)
        p_sep.paragraph_format.space_after = Pt(3)

    out_path = os.path.join(OUTPUT_DIR, "Evaluacion_Final_Ciencias_6Basico.docx")
    doc.save(out_path)
    print(f"6to Basico Test generated successfully: {out_path}")

if __name__ == "__main__":
    build_evaluacion_estudiante_6b()
