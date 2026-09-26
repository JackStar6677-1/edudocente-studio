"""
Generador de Evaluación Final y Pauta de Corrección: Orientación 5° Básico A
Colegio Luis Pasteur Anexo - Profesora Margarita Miranda B.
Año Escolar 2026
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

DOWNLOADS_ROOT = r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita"
DEST_ORIENTACION = os.path.join(DOWNLOADS_ROOT, "04 - Orientación", "5° Básico (5°A)")
os.makedirs(DEST_ORIENTACION, exist_ok=True)

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_section_header,
    set_cell_shading, set_cell_borders, set_cell_margins, COLOR_NAVY_HEX,
    COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_CORRECT_HEX,
    COLOR_NAVY_RGB, COLOR_CORRECT_RGB, COLOR_DARK_RGB, COLOR_WHITE_RGB
)

def add_orientacion_info(doc, is_pauta=False):
    t_info = doc.add_table(rows=2, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False
    w_left, w_right = Pt(255), Pt(256.2)

    filas = [
        ("Nombre del/la Estudiante: " + ("PAUTA OFICIAL DOCENTE" if is_pauta else "____________________________________"),
         "Curso: 5° Básico A"),
        ("Fecha: ______ de _________________ de 2026",
         "Puntaje Total: 25 puntos   |   Puntaje Obtenido: " + ("25 pts" if is_pauta else "_____ pts"))
    ]

    for r_idx, (d1, d2) in enumerate(filas):
        c1, c2 = t_info.cell(r_idx, 0), t_info.cell(r_idx, 1)
        c1.width, c2.width = w_left, w_right
        for c, text in [(c1, d1), (c2, d2)]:
            set_cell_margins(c, top=55, bottom=55, left=80, right=80)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            if r_idx == 0:
                set_cell_shading(c, "EAF9E6" if is_pauta else COLOR_ICE_HEX)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if is_pauta and "PAUTA" in text:
                r.bold = True
                r.font.color.rgb = COLOR_CORRECT_RGB
            elif "Puntaje" in text or "Curso" in text:
                r.bold = True

    # Cuadro de instrucciones / OA
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    t_inst = doc.add_table(rows=1, cols=1)
    t_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_inst.autofit = False
    c_ins = t_inst.cell(0, 0)
    c_ins.width = Pt(511.2)
    set_cell_shading(c_ins, COLOR_BOX_BG_HEX)
    set_cell_margins(c_ins, top=50, bottom=50, left=80, right=80)
    set_cell_borders(c_ins, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'val': 'single', 'sz': '12', 'color': COLOR_NAVY_HEX},
                            right={'color': COLOR_BORDER_HEX})
    p_ins = c_ins.paragraphs[0]
    p_ins.paragraph_format.space_before = Pt(0)
    p_ins.paragraph_format.space_after = Pt(0)

    r_oa = p_ins.add_run("Objetivo de Aprendizaje: OA 5 — Reconocer y describir las causas y consecuencias del consumo de drogas.\n")
    r_oa.bold = True
    r_oa.font.name = "Arial"
    r_oa.font.size = Pt(8.5)
    r_oa.font.color.rgb = COLOR_NAVY_RGB

    r_cont = p_ins.add_run("Contenidos: Prevención de situaciones de riesgo, factores protectores, hábitos de higiene y conductas de autocuidado.\n")
    r_cont.bold = True
    r_cont.font.name = "Arial"
    r_cont.font.size = Pt(8)

    r_txt = p_ins.add_run(
        "Instrucciones: Lee comprensivamente cada ítem y responde con letra clara y ordenada. "
        "En caso de dudas durante la evaluación, levanta tu mano en silencio para que la profesora se acerque a orientarte."
    )
    r_txt.font.name = "Arial"
    r_txt.font.size = Pt(8)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ==============================================================================
# EVALUACIÓN ESTUDIANTE
# ==============================================================================
def build_evaluacion_estudiante_orientacion():
    doc = create_base_doc()

    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="ORIENTACIÓN — 5° BÁSICO A")
    add_title_banner(doc, "EVALUACIÓN FINAL DE ORIENTACIÓN: AUTOCUIDADO Y PREVENCIÓN")
    add_orientacion_info(doc, is_pauta=False)

    # ------------------ ÍTEM I: DIBUJO Y EXPRESIÓN ------------------
    add_section_header(doc, "ÍTEM I: DIBUJO Y EXPRESIÓN DE SITUACIONES DE AUTOCUIDADO (2 pts c/u — Total: 10 pts)")

    p_i1 = doc.add_paragraph()
    p_i1.paragraph_format.space_before = Pt(1)
    p_i1.paragraph_format.space_after = Pt(4)
    r_i1 = p_i1.add_run(
        "Dibuja en cada rectángulo una situación o conducta positiva de prevención y autocuidado que practiques "
        "en el colegio o en tu hogar. Debajo de cada recuadro, escribe en la línea qué situación intentaste expresar:"
    )
    r_i1.font.name = "Arial"
    r_i1.font.size = Pt(8.5)

    # 5 recuadros de dibujo
    titulos_cajas = [
        "1. Situación de prevención en el colegio o el hogar:",
        "2. Situación de prevención en el colegio o el hogar:",
        "3. Situación de prevención en el colegio o el hogar:",
        "4. Situación de prevención en el colegio o el hogar:",
        "5. Situación de prevención en el colegio o el hogar:"
    ]

    for tit in titulos_cajas:
        t_box = doc.add_table(rows=2, cols=1)
        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_box.autofit = False

        # Celda 0: Título y Espacio para dibujar
        c_draw = t_box.cell(0, 0)
        c_draw.width = Pt(511.2)
        set_cell_shading(c_draw, "FAFCFF")
        set_cell_margins(c_draw, top=40, bottom=40, left=70, right=70)
        set_cell_borders(c_draw, top={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                 bottom={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                 left={'val': 'dashed', 'sz': '6', 'color': '888888'},
                                 right={'val': 'dashed', 'sz': '6', 'color': '888888'})
        p_d = c_draw.paragraphs[0]
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(45) # Espacio amplio para dibujar
        r_t = p_d.add_run(f"{tit}\n")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8)
        r_t.font.color.rgb = COLOR_NAVY_RGB

        r_sub = p_d.add_run("(Dibuja aquí tu situación de autocuidado)")
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(7.5)
        r_sub.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        r_sub.italic = True

        # Celda 1: Línea de escritura
        c_desc = t_box.cell(1, 0)
        c_desc.width = Pt(511.2)
        set_cell_shading(c_desc, "FFFFFF")
        set_cell_margins(c_desc, top=30, bottom=30, left=70, right=70)
        set_cell_borders(c_desc, top={'val': 'none'}, bottom={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})
        p_desc = c_desc.paragraphs[0]
        p_desc.paragraph_format.space_before = Pt(1)
        p_desc.paragraph_format.space_after = Pt(2)
        r_ex = p_desc.add_run("Descripción de lo que dibujé: ____________________________________________________________________________________")
        r_ex.font.name = "Arial"
        r_ex.font.size = Pt(8)
        r_ex.font.color.rgb = COLOR_DARK_RGB

        # Espaciador entre cajas
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(2)

    # ------------------ ÍTEM II: ESCRITURA DE HIGIENE PERSONAL ------------------
    add_section_header(doc, "ÍTEM II: HÁBITOS DE CUIDADO E HIGIENE PERSONAL (1 pto c/u — Total: 5 pts)")

    p_i2 = doc.add_paragraph()
    p_i2.paragraph_format.space_before = Pt(1)
    p_i2.paragraph_format.space_after = Pt(4)
    r_i2 = p_i2.add_run("Escribe cinco formas o hábitos de cuidado e higiene personal en el hogar y en el colegio que ayuden a mantener tu salud:")
    r_i2.font.name = "Arial"
    r_i2.font.size = Pt(8.5)

    t_lines = doc.add_table(rows=5, cols=2)
    t_lines.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_lines.autofit = False

    for num in range(1, 6):
        c_n = t_lines.cell(num - 1, 0)
        c_l = t_lines.cell(num - 1, 1)
        c_n.width, c_l.width = Pt(25), Pt(486.2)

        set_cell_shading(c_n, COLOR_ICE_HEX)
        set_cell_margins(c_n, top=40, bottom=40, left=20, right=20)
        set_cell_borders(c_n, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                              left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
        p_n = c_n.paragraphs[0]
        p_n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_n.paragraph_format.space_before = Pt(0)
        p_n.paragraph_format.space_after = Pt(0)
        rn = p_n.add_run(f"{num}.")
        rn.bold = True
        rn.font.name = "Arial"
        rn.font.size = Pt(8.5)
        rn.font.color.rgb = COLOR_NAVY_RGB

        set_cell_shading(c_l, "FFFFFF")
        set_cell_margins(c_l, top=40, bottom=40, left=40, right=40)
        set_cell_borders(c_l, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                              left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
        p_l = c_l.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(0)
        p_l.paragraph_format.space_after = Pt(0)
        rl = p_l.add_run("______________________________________________________________________________________________")
        rl.font.name = "Arial"
        rl.font.size = Pt(8)
        rl.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # ------------------ ÍTEM III: VERDADERO O FALSO ------------------
    add_section_header(doc, "ÍTEM III: VERDADERO O FALSO (1 pto c/u — Total: 10 pts)")

    p_i3 = doc.add_paragraph()
    p_i3.paragraph_format.space_before = Pt(1)
    p_i3.paragraph_format.space_after = Pt(4)
    r_i3 = p_i3.add_run(
        "Lee atentamente cada oración y escribe en el paréntesis una "
    )
    r_i3.font.name = "Arial"
    r_i3.font.size = Pt(8.5)
    rv = p_i3.add_run("V")
    rv.bold = True
    rv.font.name = "Arial"
    rv.font.size = Pt(8.5)
    r_i3b = p_i3.add_run(" si es verdadera o una ")
    r_i3b.font.name = "Arial"
    r_i3b.font.size = Pt(8.5)
    rf = p_i3.add_run("F")
    rf.bold = True
    rf.font.name = "Arial"
    rf.font.size = Pt(8.5)
    r_i3c = p_i3.add_run(" si es falsa:")
    r_i3c.font.name = "Arial"
    r_i3c.font.size = Pt(8.5)

    afirmaciones = [
        "En el colegio debes bajar las escaleras corriendo.",
        "Debes mantener una alimentación saludable con productos que tengan tres sellos.",
        "Al toque de timbre sales corriendo de la sala de clases.",
        "Después de actividades físicas debes lavar tus manos y tu cara.",
        "Lavarse las manos con agua y jabón antes de comer en el recreo previene enfermedades.",
        "Correr a máxima velocidad por pasillos llenos de gente ayuda a evitar accidentes.",
        "Avisa a un profesor o adulto del colegio si sufres una herida o malestar en conducta de autocuidado.",
        "Cruzar el patio sin mirar evita cualquier tipo de lesión.",
        "Utilizar tijeras escolares solo para tareas indicadas evita cortes y accidentes.",
        "Participar en clases en forma respetuosa levantando mi mano para dar mi opinión."
    ]

    t_vf = doc.add_table(rows=len(afirmaciones), cols=3)
    t_vf.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_vf.autofit = False

    for idx, texto in enumerate(afirmaciones, 1):
        c_num = t_vf.cell(idx - 1, 0)
        c_box = t_vf.cell(idx - 1, 1)
        c_txt = t_vf.cell(idx - 1, 2)

        c_num.width = Pt(22)
        c_box.width = Pt(38)
        c_txt.width = Pt(451.2)

        bg_col = "FFFFFF" if idx % 2 != 0 else COLOR_BOX_BG_HEX

        for c in (c_num, c_box, c_txt):
            set_cell_shading(c, bg_col)
            set_cell_margins(c, top=35, bottom=35, left=35, right=35)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

        p_n = c_num.paragraphs[0]
        p_n.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_n.paragraph_format.space_before = Pt(0)
        p_n.paragraph_format.space_after = Pt(0)
        rn = p_n.add_run(f"{idx}.")
        rn.font.name = "Arial"
        rn.font.size = Pt(8.5)
        rn.bold = True
        rn.font.color.rgb = COLOR_NAVY_RGB

        p_b = c_box.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(0)
        rb = p_b.add_run("(     )")
        rb.font.name = "Arial"
        rb.font.size = Pt(8.5)
        rb.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        p_t = c_txt.paragraphs[0]
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(0)
        rt = p_t.add_run(texto)
        rt.font.name = "Arial"
        rt.font.size = Pt(8.5)

    out_name = "Evaluacion_Final_Orientacion_5Basico.docx"
    doc.save(os.path.join(OUTPUT_DIR, out_name))
    doc.save(os.path.join(DEST_ORIENTACION, out_name))
    print(f"  [OK] Prueba Estudiante guardada: {out_name}")

# ==============================================================================
# PAUTA DE CORRECCIÓN (DOCENTE)
# ==============================================================================
def build_pauta_correccion_orientacion():
    doc = create_base_doc()

    add_header(doc, school="COLEGIO LUIS PASTEUR ANEXO", subject="ORIENTACIÓN — 5° BÁSICO A", subtitle="DOCUMENTO DOCENTE")
    add_title_banner(doc, "PAUTA DE CORRECCIÓN: EVALUACIÓN FINAL DE ORIENTACIÓN", is_pauta=True)
    add_orientacion_info(doc, is_pauta=True)

    # ------------------ SOLUCIONARIO ÍTEM I ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM I: DIBUJO Y EXPRESIÓN DE AUTOCUIDADO (10 pts)")

    p_i1 = doc.add_paragraph()
    p_i1.paragraph_format.space_before = Pt(1)
    p_i1.paragraph_format.space_after = Pt(4)
    r_i1 = p_i1.add_run("Criterio de Corrección (2 puntos por recuadro: 1 punto por dibujo pertinente + 1 punto por descripción clara escrita por el estudiante).")
    r_i1.bold = True
    r_i1.font.name = "Arial"
    r_i1.font.size = Pt(8.5)
    r_i1.font.color.rgb = COLOR_CORRECT_RGB

    respuestas_item1 = [
        ("1. No correr en las escaleras",
         "Ejemplo esperado de dibujo: Estudiante caminando de forma tranquila y sosteniéndose del pasamanos.\n"
         "Justificación pedagógica: Evita caídas graves, resbalones y fracturas. Conducta primordial de seguridad escolar."),

        ("2. No empujar a compañeros",
         "Ejemplo esperado de dibujo: Estudiantes jugando respetuosamente o haciendo fila sin empujones.\n"
         "Justificación pedagógica: Fomenta la sana convivencia y previene golpes contra el suelo o mobiliario."),

        ("3. Lavado de manos y cara",
         "Ejemplo esperado de dibujo: Estudiante en el lavamanos usando agua y jabón antes de comer o tras jugar.\n"
         "Justificación pedagógica: Hábito esencial de higiene que elimina virus y bacterias transmisoras de infecciones."),

        ("4. Comer saludable / Colaciones nutritivas",
         "Ejemplo esperado de dibujo: Estudiante comiendo frutas, frutos secos o bebiendo agua.\n"
         "Justificación pedagógica: Proporciona nutrientes esenciales y protege el desarrollo físico e intelectual."),

        ("5. No usar el celular en clases",
         "Ejemplo esperado de dibujo: Celular guardado en la mochila y estudiante atento a la pizarra.\n"
         "Justificación pedagógica: Protege la atención, el respeto docente y evita distracciones que perjudican el aprendizaje.")
    ]

    for tit, detalle in respuestas_item1:
        t_box = doc.add_table(rows=1, cols=1)
        t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_box.autofit = False
        c = t_box.cell(0, 0)
        c.width = Pt(511.2)
        set_cell_shading(c, "F6FAFE")
        set_cell_margins(c, top=45, bottom=45, left=70, right=70)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'val': 'single', 'sz': '12', 'color': COLOR_CORRECT_HEX},
                            right={'color': COLOR_BORDER_HEX})
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        rt = p.add_run(f"• Situación esperada: {tit}\n")
        rt.bold = True
        rt.font.name = "Arial"
        rt.font.size = Pt(8.5)
        rt.font.color.rgb = COLOR_CORRECT_RGB

        rd = p.add_run(detalle)
        rd.font.name = "Arial"
        rd.font.size = Pt(8)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ------------------ SOLUCIONARIO ÍTEM II ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM II: FORMAS DE HIGIENE PERSONAL (5 pts)")

    respuestas_item2 = [
        "1. Usar mascarilla si estás resfriado (evita la dispersión de gotas y contagio de virus a compañeros).",
        "2. Usa toalla y artículos de aseo en la asignatura de educación física (mantiene la higiene personal tras la transpiración).",
        "3. No tirar basura en el suelo (mantiene un ambiente escolar limpio, libre de gérmenes y agradable para todos).",
        "4. Cuidar las áreas verdes del colegio (promueve un entorno saludable y el respeto por el medio ambiente escolar).",
        "5. Mantener ordenado su espacio personal (facilita la limpieza del pupitre, la comodidad y previene extravíos de útiles)."
    ]

    t_r2 = doc.add_table(rows=len(respuestas_item2), cols=1)
    t_r2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_r2.autofit = False

    for idx, r_txt in enumerate(respuestas_item2):
        c = t_r2.cell(idx, 0)
        c.width = Pt(511.2)
        set_cell_shading(c, "FFFFFF" if idx % 2 == 0 else COLOR_BOX_BG_HEX)
        set_cell_margins(c, top=35, bottom=35, left=60, right=60)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"• {r_txt}")
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        if idx in [0, 1]:
            r.bold = True
            r.font.color.rgb = COLOR_CORRECT_RGB

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # ------------------ SOLUCIONARIO ÍTEM III ------------------
    add_section_header(doc, "SOLUCIONARIO ÍTEM III: VERDADERO O FALSO (10 pts)")

    pautas_vf = [
        ("F", "En el colegio debes bajar las escaleras corriendo.",
         "Es falsa. Correr en escaleras es de alto riesgo de caídas graves; siempre se debe caminar y tomarse del pasamanos."),

        ("F", "Debes mantener una alimentación saludable con productos que tengan tres sellos.",
         "Es falsa. Los sellos de advertencia ('Alto en azúcares, grasas, sodio o calorías') indican exceso perjudicial; lo sano es preferir productos sin sellos o naturales."),

        ("F", "Al toque de timbre sales corriendo de la sala de clases.",
         "Es falsa. Salir corriendo en masa genera aglomeraciones y caídas en la puerta; se debe salir de forma ordenada y caminando."),

        ("V", "Después de actividades físicas debes lavar tus manos y tu cara.",
         "Es verdadera. Remueve el sudor, refresca y previene infecciones dermatológicas y proliferación bacteriana."),

        ("V", "Lavarse las manos con agua y jabón antes de comer en el recreo previene enfermedades.",
         "Es verdadera. La higiene de manos antes de ingerir alimentos es la principal barrera contra bacterias estomacales."),

        ("F", "Correr a máxima velocidad por pasillos llenos de gente ayuda a evitar accidentes.",
         "Es falsa. Correr en pasillos concurridos es causa frecuente de colisiones y lesiones; se debe transitar caminando."),

        ("V", "Avisa a un profesor o adulto del colegio si sufres una herida o malestar en conducta de autocuidado.",
         "Es verdadera. La búsqueda de ayuda con un adulto responsable es la base del autocuidado y primeros auxilios oportunos."),

        ("F", "Cruzar el patio sin mirar evita cualquier tipo de lesión.",
         "Es falsa. Cruzar sin mirar expone a recibir balonazos o chocar con otros estudiantes que están jugando."),

        ("V", "Utilizar tijeras escolares solo para tareas indicadas evita cortes y accidentes.",
         "Es verdadera. Las tijeras deben usarse exclusivamente para cortar materiales escolares de manera segura sobre la mesa."),

        ("V", "Participar en clases en forma respetuosa levantando mi mano para dar mi opinión.",
         "Es verdadera. Favorece la buena convivencia escolar, la escucha activa y el clima positivo de aprendizaje en el aula.")
    ]

    t_vf_sol = doc.add_table(rows=len(pautas_vf) + 1, cols=4)
    t_vf_sol.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_vf_sol.autofit = False

    w_cols = [Pt(25), Pt(45), Pt(230), Pt(211.2)]
    headers_vf = ["N°", "Resp.", "Afirmación Evaluada", "Justificación Pedagógica"]

    for i, h in enumerate(headers_vf):
        c = t_vf_sol.cell(0, i)
        c.width = w_cols[i]
        set_cell_shading(c, COLOR_NAVY_HEX)
        set_cell_margins(c, top=45, bottom=45, left=45, right=45)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = COLOR_WHITE_RGB

    for row_idx, (letra, afirmacion, justificacion) in enumerate(pautas_vf, 1):
        bg = "FFFFFF" if row_idx % 2 != 0 else COLOR_BOX_BG_HEX
        c_n = t_vf_sol.cell(row_idx, 0)
        c_l = t_vf_sol.cell(row_idx, 1)
        c_a = t_vf_sol.cell(row_idx, 2)
        c_j = t_vf_sol.cell(row_idx, 3)

        for col_idx, cell in enumerate((c_n, c_l, c_a, c_j)):
            cell.width = w_cols[col_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=35, bottom=35, left=40, right=40)
            set_cell_borders(cell, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                   left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})

        pn = c_n.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pn.paragraph_format.space_before = Pt(0)
        pn.paragraph_format.space_after = Pt(0)
        rn = pn.add_run(f"{row_idx}")
        rn.font.name = "Arial"
        rn.font.size = Pt(8)
        rn.bold = True

        pl = c_l.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pl.paragraph_format.space_before = Pt(0)
        pl.paragraph_format.space_after = Pt(0)
        rl = pl.add_run(f"[ {letra} ]")
        rl.font.name = "Arial"
        rl.font.size = Pt(8.5)
        rl.bold = True
        rl.font.color.rgb = COLOR_CORRECT_RGB

        pa = c_a.paragraphs[0]
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after = Pt(0)
        ra = pa.add_run(afirmacion)
        ra.font.name = "Arial"
        ra.font.size = Pt(8)

        pj = c_j.paragraphs[0]
        pj.paragraph_format.space_before = Pt(0)
        pj.paragraph_format.space_after = Pt(0)
        rj = pj.add_run(justificacion)
        rj.font.name = "Arial"
        rj.font.size = Pt(7.5)
        rj.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    out_name = "Pauta_Correccion_Orientacion_5Basico.docx"
    doc.save(os.path.join(OUTPUT_DIR, out_name))
    doc.save(os.path.join(DEST_ORIENTACION, out_name))
    print(f"  [OK] Pauta Docente guardada: {out_name}")

def generar_orientacion_completa():
    print("\n=======================================================")
    print(" GENERANDO EVALUACIÓN Y PAUTA DE ORIENTACIÓN 5° BÁSICO")
    print("=======================================================")
    build_evaluacion_estudiante_orientacion()
    build_pauta_correccion_orientacion()
    print(f"\n[COMPLETADO] Material de Orientación 5° Básico desplegado en: {DEST_ORIENTACION}\n")

if __name__ == "__main__":
    generar_orientacion_completa()
