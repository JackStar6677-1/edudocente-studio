import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Entrevistas Apoderados"

# Encabezado institucional
ws.merge_cells("A1:G1")
ws["A1"] = "COLEGIO CASTELGANDOLFO - CRONOGRAMA DE ENTREVISTAS DE APODERADOS 2026"
ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 28

ws.merge_cells("A2:G2")
ws["A2"] = "Jefatura de Curso y Docencia Institucional"
ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="1E3A5F")
ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

# Encabezados de tabla
headers = [
    "N°", "Nombre del apoderado/a", "Correo electrónico",
    "Nombre del/la estudiante", "Curso", "Fecha entrevista", "Hora de entrevista"
]

ws.row_dimensions[4].height = 25
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(4, col_idx, h)
    cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

data = []

thin_border = Border(
    left=Side(style='thin', color='B0C4DE'),
    right=Side(style='thin', color='B0C4DE'),
    top=Side(style='thin', color='B0C4DE'),
    bottom=Side(style='thin', color='B0C4DE')
)

row_start = 5
for r_data in data:
    ws.row_dimensions[row_start].height = 20
    is_grupo1 = "07 de octubre" in r_data[5]
    fill_bg = "F6FAFE" if is_grupo1 else "FFFFFF"

    for c_i, val in enumerate(r_data, 1):
        cell = ws.cell(row_start, c_i, val)
        cell.font = Font(name="Arial", size=9)
        cell.border = thin_border
        cell.fill = PatternFill(start_color=fill_bg, end_color=fill_bg, fill_type="solid")
        if c_i in [1, 5, 6, 7]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
    row_start += 1

def generar_cronograma_excel():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, "output")
    os.makedirs(output_dir, exist_ok=True)

    # Ajustar anchos de columnas
    col_widths = {1: 6, 2: 34, 3: 38, 4: 36, 5: 14, 6: 24, 7: 20}
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    local_path = os.path.join(output_dir, "Cronograma_Entrevistas_Apoderados_2026.xlsx")
    wb.save(local_path)
    print(f"Excel guardado localmente en: {local_path}")

    # Si existe la carpeta de descargas de la profesora, guardar copia también
    downloads_path = r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita\03 - Registro y Gestión Docente\Nóminas y Listas"
    if os.path.exists(downloads_path):
        target_path = os.path.join(downloads_path, "Cronograma_Entrevistas_Apoderados_2026.xlsx")
        try:
            wb.save(target_path)
            print(f"Copia actualizada en Descargas: {target_path}")
        except Exception as e:
            print(f"Aviso al guardar en Descargas: {e}")

    return local_path

if __name__ == "__main__":
    generar_cronograma_excel()
