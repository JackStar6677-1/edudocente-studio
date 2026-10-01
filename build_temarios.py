"""
Generador Automatizado de Temarios e Informativos por Curso
Colegio Castelgandolfo - Docencia Institucional
Año Escolar 2026
"""

import os
import sys
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_TEMARIOS_DIR = os.path.join(BASE_DIR, "output", "temarios")
os.makedirs(OUTPUT_TEMARIOS_DIR, exist_ok=True)

DOWNLOADS_TEMARIOS_DIR = r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita\00 - Temarios por Curso (Para Enviar a Apoderados)"
os.makedirs(DOWNLOADS_TEMARIOS_DIR, exist_ok=True)

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_section_header,
    set_cell_shading, set_cell_borders, set_cell_margins, COLOR_NAVY_HEX,
    COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX, COLOR_NAVY_RGB,
    COLOR_DARK_RGB, COLOR_WHITE_RGB, safe_save
)

def build_temario_base(curso_header, subject_header, title_banner, datos_caja, saludo_parrafo, secciones_datos, orientaciones_estudio, nota_extra=None):
    doc = create_base_doc()

    # Encabezado institucional con logo
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject=f"{subject_header} — {curso_header}")
    add_title_banner(doc, title_banner)

    # Cuadro informativo de datos
    t_info = doc.add_table(rows=len(datos_caja), cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False
    w_left, w_right = Pt(255), Pt(256.2)

    for r_idx, (d1, d2) in enumerate(datos_caja):
        c1, c2 = t_info.cell(r_idx, 0), t_info.cell(r_idx, 1)
        c1.width, c2.width = w_left, w_right
        for c, text in [(c1, d1), (c2, d2)]:
            set_cell_margins(c, top=60, bottom=60, left=90, right=90)
            set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            if r_idx == 0:
                set_cell_shading(c, COLOR_BOX_BG_HEX)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if "Estimados" in text or "Asignatura" in text or "Docente" in text or "Objetivo" in text:
                r.bold = True

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(3)

    # Saludo y Presentación del Correo
    p_saludo = doc.add_paragraph()
    p_saludo.paragraph_format.space_before = Pt(2)
    p_saludo.paragraph_format.space_after = Pt(4)
    r_sal = p_saludo.add_run(saludo_parrafo)
    r_sal.font.name = "Arial"
    r_sal.font.size = Pt(8.5)

    # Sección 1: Contenidos Temáticos
    add_section_header(doc, "1. CONTENIDOS TEMÁTICOS Y REFERENCIAS")

    headers_tabla = secciones_datos["headers"]
    filas_tabla = secciones_datos["filas"]
    widths_tabla = secciones_datos["widths"]

    t_temas = doc.add_table(rows=len(filas_tabla) + 1, cols=len(headers_tabla))
    t_temas.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_temas.autofit = False

    # Cabecera de la tabla
    for i, h in enumerate(headers_tabla):
        c = t_temas.cell(0, i)
        c.width = widths_tabla[i]
        set_cell_shading(c, COLOR_NAVY_HEX)
        set_cell_margins(c, top=60, bottom=60, left=70, right=70)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = COLOR_WHITE_RGB

    # Filas de contenidos
    for row_idx, data_row in enumerate(filas_tabla, 1):
        bg_col = "FFFFFF" if row_idx % 2 != 0 else COLOR_BOX_BG_HEX
        for col_idx, text_content in enumerate(data_row):
            cell = t_temas.cell(row_idx, col_idx)
            cell.width = widths_tabla[col_idx]
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            set_cell_borders(cell, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                                   left={'color': COLOR_BORDER_HEX}, right={'color': COLOR_BORDER_HEX})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if col_idx == 1 and len(headers_tabla) == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            r = p.add_run(text_content)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if col_idx == 0:
                r.bold = True

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(0)
    p_sp2.paragraph_format.space_after = Pt(3)

    # Sección 2: Orientaciones Pedagógicas y de Estudio
    add_section_header(doc, "2. ORIENTACIONES PEDAGÓGICAS PARA EL ESTUDIO EN EL HOGAR")

    t_tips = doc.add_table(rows=len(orientaciones_estudio), cols=1)
    t_tips.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_tips.autofit = False

    for idx, (title_tip, desc_tip) in enumerate(orientaciones_estudio):
        c = t_tips.cell(idx, 0)
        c.width = Pt(511.2)
        set_cell_shading(c, COLOR_BOX_BG_HEX)
        set_cell_margins(c, top=45, bottom=45, left=80, right=80)
        set_cell_borders(c, top={'color': COLOR_BORDER_HEX}, bottom={'color': COLOR_BORDER_HEX},
                            left={'val': 'single', 'sz': '12', 'color': COLOR_NAVY_HEX},
                            right={'color': COLOR_BORDER_HEX})
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r_t = p.add_run(f"• {title_tip}: ")
        r_t.bold = True
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.5)
        r_t.font.color.rgb = COLOR_NAVY_RGB

        r_d = p.add_run(desc_tip)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8.5)
        r_d.font.color.rgb = COLOR_DARK_RGB

    if nota_extra:
        p_extra = doc.add_paragraph()
        p_extra.paragraph_format.space_before = Pt(3)
        p_extra.paragraph_format.space_after = Pt(2)
        r_ex = p_extra.add_run(f"Nota importante: {nota_extra}")
        r_ex.font.name = "Arial"
        r_ex.font.size = Pt(8)
        r_ex.italic = True
        r_ex.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Firma institucional
    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(0)
    p_sp3.paragraph_format.space_after = Pt(2)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(2)
    p_sig.paragraph_format.space_after = Pt(0)

    sig_runs = [
        ("Saludos cordiales,\n", False, 8.5),
        ("Docencia Institucional\n", True, 9),
        ("Docente de Educación Básica | Colegio Castelgandolfo\n", False, 8),
        ("Contacto: ", False, 8),
        ("docencia@colegiocastelgandolfo.cl", True, 8)
    ]
    for text, bold, sz in sig_runs:
        r = p_sig.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(sz)
        r.bold = bold
        r.font.color.rgb = COLOR_NAVY_RGB if bold else COLOR_DARK_RGB

    return doc

# ==============================================================================
# 1. TEMARIO CIENCIAS NATURALES - 5° BÁSICO A Y B
# ==============================================================================
def generar_temario_5ciencias():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (5° Básico A y B)", "Asignatura: Ciencias Naturales"),
        ("Objetivo de Aprendizaje: CN05 OA 11", "Evaluación Final Año Escolar 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Ciencias Naturales para el estudio de la evaluación final año 2026. A continuación, "
        "se detallan las páginas del texto escolar Mineduc y los temas prioritarios a repasar con dedicación:"
    )
    secciones_datos = {
        "headers": ["Contenido Específico", "Páginas del Texto", "Conceptos Clave a Reforzar"],
        "widths": [Pt(145), Pt(105), Pt(261.2)],
        "filas": [
            ("¿Qué es la energía eléctrica?", "Página 146",
             "Definición de energía eléctrica. Manifestaciones en la vida cotidiana (luz, calor, sonido, movimiento). Transformación de la energía en aparatos del hogar e importancia del uso responsable y ahorro energético."),
            ("Circuitos eléctricos", "Página 149",
             "Definición y funcionamiento de un circuito eléctrico cerrado y abierto. Función de sus partes: fuente de energía, interruptor, receptores (ampolletas) y cables conductores."),
            ("Símbolos de un circuito eléctrico", "Página 150",
             "Reconocimiento y dibujo de la simbología normalizada: fuente/pila (+ y –), cable conductor, interruptor (abierto/cerrado) y receptor/ampolleta (círculo con X)."),
            ("Conductores y aisladores eléctricos", "Página 152",
             "Diferenciación entre materiales conductores de electricidad (metales, agua con sales) y materiales aisladores (plásticos, goma, madera seca, vidrio). Normas de seguridad eléctrica en el hogar.")
        ]
    }
    orientaciones = [
        ("Lectura comprensiva del texto escolar", "Leer detenidamente las páginas 146, 149, 150 y 152 del libro del estudiante, subrayando ideas principales y glosario."),
        ("Práctica de dibujo de símbolos", "Ejercitar en el cuaderno el dibujo claro de los 4 símbolos eléctricos (pila, ampolleta, interruptor y cable)."),
        ("Revisión de actividades del cuaderno", "Repasar las actividades de circuitos y conductores realizadas durante las clases presenciales.")
    ]
    doc = build_temario_base(
        curso_header="5° BÁSICO A Y B",
        subject_header="CIENCIAS NATURALES",
        title_banner="TEMARIO Y CONTENIDOS EVALUACIÓN FINAL: ENERGÍA ELÉCTRICA Y CIRCUITOS",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones
    )
    out_name = "Temario_Ciencias_5Basico_A_B.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

# ==============================================================================
# 2. TEMARIO CIENCIAS NATURALES - 6° BÁSICO A Y B
# ==============================================================================
def generar_temario_6ciencias():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (6° Básico A y B)", "Asignatura: Ciencias Naturales"),
        ("Objetivo de Aprendizaje: CN06 OA 13", "Evaluación Final Año Escolar 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Ciencias Naturales para el estudio de la evaluación final año 2026. A continuación, "
        "se presenta el temario debidamente explicado para guiar el estudio de los cambios de estado de la materia:"
    )
    secciones_datos = {
        "headers": ["Contenido Explicado", "Páginas del Texto", "Detalle Pedagógico y Procesos Físicos"],
        "widths": [Pt(150), Pt(100), Pt(261.2)],
        "filas": [
            ("Cambios de estado de la materia", "Páginas 170 y 173",
             "Comprender que son transformaciones físicas: la materia solo cambia su aspecto externo, volumen y orden de partículas, pero sigue siendo la misma sustancia química. Organización molecular en sólidos, líquidos y gases."),
            ("Cambios de estado por absorción de calor (Progresivos)", "Páginas 173 a 175",
             "Al absorber calor del entorno, las partículas aumentan su energía cinética (se mueven más rápido) y se separan venciendo la fuerza de atracción:\n"
             "• Fusión: de sólido a líquido.\n"
             "• Vaporización (Evaporación y Ebullición): de líquido a gas.\n"
             "• Sublimación progresiva: de sólido directo a gas sin pasar por líquido (ej. hielo seco, naftalina)."),
            ("Cambios de estado por liberación de calor (Regresivos)", "Páginas 181 a 183",
             "Al liberar calor (enfriamiento), las partículas pierden energía cinética, se mueven más lento y las fuerzas de atracción las acercan y ordenan:\n"
             "• Condensación: de gas a líquido (ej. vapor sobre ventana fría, rocío).\n"
             "• Solidificación: de líquido a sólido (ej. agua a hielo).\n"
             "• Sublimación regresiva / inversa: de gas directo a sólido (ej. escarcha)."),
            ("Modelo corpuscular de partículas", "Páginas 170 y 173",
             "Identificar y representar el comportamiento de las partículas: ordenadas y compactas en estado sólido; unidas pero con movilidad en estado líquido; muy separadas y veloces en estado gaseoso.")
        ]
    }
    orientaciones = [
        ("Distinción entre absorción y liberación", "Asegurarse de comprender qué procesos requieren suministrar calor y cuáles ocurren al enfriarse o perder calor."),
        ("Ejemplos cotidianos del ciclo del agua", "Relacionar cada cambio de estado con vivencias prácticas en la cocina, el clima y la naturaleza."),
        ("Dibujo y esquemas de partículas", "Practicar el dibujo de partículas en los tres estados fundamentales para reforzar la comprensión visual.")
    ]
    doc = build_temario_base(
        curso_header="6° BÁSICO A Y B",
        subject_header="CIENCIAS NATURALES",
        title_banner="TEMARIO Y CONTENIDOS EVALUACIÓN FINAL: CAMBIOS DE ESTADO DE LA MATERIA",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones
    )
    out_name = "Temario_Ciencias_6Basico_A_B.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

# ==============================================================================
# 3. TEMARIO CIENCIAS NATURALES - 8° BÁSICO A
# ==============================================================================
def generar_temario_8ciencias():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (8° Básico A)", "Asignatura: Ciencias Naturales"),
        ("Objetivo de Aprendizaje: CN08 OA 14", "Evaluación Final Año Escolar 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Ciencias Naturales para el estudio de la evaluación final año 2026. A continuación, "
        "se detallan los contenidos sobre la estructura atómica de la materia y los modelos históricos:"
    )
    secciones_datos = {
        "headers": ["Contenido Específico", "Páginas del Texto", "Conceptos Clave y Desarrollo Científico"],
        "widths": [Pt(150), Pt(100), Pt(261.2)],
        "filas": [
            ("Teoría Atómica de John Dalton", "Página 126",
             "Primer modelo científico formal (1808). Postulados fundamentales: la materia está formada por partículas indivisibles llamadas átomos; átomos de un mismo elemento son idénticos; formación de compuestos en proporciones fijas; conservación del átomo en reacciones químicas."),
            ("Evolución de los Modelos Atómicos", "Páginas 127 a 133",
             "Aportes de la investigación experimental a lo largo de la historia:\n"
             "• J.J. Thomson: descubrimiento del electrón y modelo de 'budín de pasas'.\n"
             "• E. Rutherford: experimento de la lámina de oro, descubrimiento del núcleo denso y positivo, protones y corteza.\n"
             "• N. Bohr: órbitas cuantizadas de energía donde giran los electrones."),
            ("Estructura del Átomo y Partículas Subatómicas", "Páginas 134 a 136",
             "Ubicación, carga y masa relativa de:\n"
             "• Protones (p+): carga positiva, ubicados en el núcleo.\n"
             "• Neutrones (n0): carga neutra, ubicados en el núcleo.\n"
             "• Electrones (e–): carga negativa, ubicados en la corteza o nube electrónica."),
            ("Número Atómico (Z) y Número Másico (A)", "Página 136",
             "Cálculo de partículas en átomos neutros:\n"
             "• Número Atómico (Z) = Número de protones (identidad del elemento).\n"
             "• Número Másico (A) = Protones + Neutrones (A = Z + n0).\n"
             "• Cálculo de neutrones: n0 = A – Z.")
        ]
    }
    orientaciones = [
        ("Línea de tiempo de modelos atómicos", "Construir en el cuaderno un esquema comparativo con las características de los modelos de Dalton, Thomson, Rutherford y Bohr."),
        ("Ejercicios de cálculo de Z y A", "Practicar la determinación de protones, electrones y neutrones utilizando la simbología estándar de la tabla periódica."),
        ("Lectura activa del texto escolar", "Revisar los experimentos históricos explicados en las páginas 126 a 136 del texto del estudiante.")
    ]
    doc = build_temario_base(
        curso_header="8° BÁSICO A",
        subject_header="CIENCIAS NATURALES",
        title_banner="TEMARIO Y CONTENIDOS EVALUACIÓN FINAL: TEORÍA ATÓMICA Y EL ÁTOMO",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones
    )
    out_name = "Temario_Ciencias_8Basico_A.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

# ==============================================================================
# 4. TEMARIO MÚSICA - 1° BÁSICO A
# ==============================================================================
def generar_temario_1musica():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (1° Básico A)", "Asignatura: Música"),
        ("Objetivos de Aprendizaje: MU01 OA 02 / OA 04", "Evaluación Final Práctica Año 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Música para la evaluación final práctica año 2026. La evaluación consistirá en "
        "la interpretación musical guiada de la canción 'Estrellita', integrando canto y ejecución instrumental:"
    )
    secciones_datos = {
        "headers": ["Área de Evaluación", "Modalidad / Repertorio", "Criterios e Indicadores de Logro"],
        "widths": [Pt(145), Pt(125), Pt(241.2)],
        "filas": [
            ("Expresión Vocal (Canto al unísono)", "Canción: 'Estrellita'\n(Estrellita, ¿dónde estás?)",
             "Cantar al unísono junto al grupo curso con afinación natural, dicción clara de las palabras, respiración adecuada y entusiasmo respetuoso."),
            ("Ejecución Instrumental a Elección", "Instrumento a elección:\n• Sonaja (Percusión)\n• Metalófono (Melódico)",
             "El estudiante seleccionará uno de los instrumentos trabajados en el semestre:\n"
             "• Opción Sonaja: Marcar con precisión el pulso rítmico constante de la canción.\n"
             "• Opción Metalófono: Tocar las notas o placas básicas trabajadas en clase con técnica adecuada de baquetas."),
            ("Habilidades y Postura Musical", "Trabajo en el aula de clases",
             "Demostrar postura corporal atenta, cuidado y respeto en el uso del instrumento musical, e inicio y cierre coordinado junto a la profesora.")
        ]
    }
    orientaciones = [
        ("Cantar la canción en familia", "Repasar la letra de 'Estrellita' en el hogar para afianzar la memoria, seguridad y pronunciación de los niños."),
        ("Práctica del pulso con palmas o instrumento", "Marcar el ritmo regular de la canción aplaudiendo o usando la sonaja/metalófono al compás de la voz."),
        ("Refuerzo positivo", "Felicitar y motivar a los pequeños reconociendo su dedicación y disfrute por la música.")
    ]
    doc = build_temario_base(
        curso_header="1° BÁSICO A",
        subject_header="MÚSICA",
        title_banner="TEMARIO Y PAUTA INFORMATIVA: EVALUACIÓN PRÁCTICA DE MÚSICA",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones,
        nota_extra="El instrumento (sonaja o metalófono) será el que el estudiante ha venido utilizando en sus clases habituales."
    )
    out_name = "Temario_Musica_1Basico_A.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

# ==============================================================================
# 5. TEMARIO MÚSICA - 2° BÁSICO A
# ==============================================================================
def generar_temario_2musica():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (2° Básico A)", "Asignatura: Música"),
        ("Objetivos de Aprendizaje: MU02 OA 02 / OA 04", "Evaluación Final Práctica Año 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Música para la evaluación final práctica año 2026. A continuación, se detallan "
        "los aspectos evaluativos de la interpretación coral e instrumental de la canción 'Estrellita':"
    )
    secciones_datos = {
        "headers": ["Eje de Aprendizaje", "Repertorio e Instrumentación", "Descriptores de Desempeño Musical"],
        "widths": [Pt(145), Pt(125), Pt(241.2)],
        "filas": [
            ("Interpretación Coral al Unísono", "Canción: 'Estrellita'\n(Repertorio infantil tradicional)",
             "Canto grupal al unísono manteniendo la afinación, coordinación rítmica con el texto lírico, claridad de pronunciación y control de dinámicas sonoras (suave/fuerte)."),
            ("Acompañamiento Instrumental Diferenciado", "Instrumento a elección del estudiante:\n• Metalófono\n• Sonaja / Percusión menor",
             "Cada estudiante escoge su instrumento de ejecución:\n"
             "• Metalófono: Interpretación melódica de la frase musical (Do-Do-Sol-Sol-La-La-Sol...), ubicando correctamente las notas en las placas y alternando baquetas.\n"
             "• Sonaja: Mantener el pulso rítmico sin adelantarse ni atrasarse, sincronizando con el canto grupal."),
            ("Expresividad y Participación", "Desempeño en conjunto",
             "Concentración, postura ergonómica adecuada, seguimiento de las entradas y cortes de la profesora, y valoración del trabajo armónico en equipo.")
        ]
    }
    orientaciones = [
        ("Práctica melódica en metalófono", "Recomendar que los estudiantes que elijan metalófono canten la nota al mismo tiempo que la tocan para reforzar el oído musical."),
        ("Seguimiento del pulso rítmico", "Practicar con metrónomo lento o palmadas para asegurar que el ritmo de la sonaja sea constante."),
        ("Escucha atenta y concentrada", "Fomentar la escucha activa de la melodía para entrar coordinados desde el inicio de la canción.")
    ]
    doc = build_temario_base(
        curso_header="2° BÁSICO A",
        subject_header="MÚSICA",
        title_banner="TEMARIO Y PAUTA INFORMATIVA: EVALUACIÓN PRÁCTICA DE MÚSICA",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones,
        nota_extra="Se evaluará en clases presenciales según el instrumento escogido por cada estudiante."
    )
    out_name = "Temario_Musica_2Basico_A.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

# ==============================================================================
# 6. TEMARIO ORIENTACIÓN - 5° BÁSICO A
# ==============================================================================
def generar_temario_5orientacion():
    datos_caja = [
        ("Estimados/as apoderados y estudiantes (5° Básico A)", "Asignatura: Orientación"),
        ("Objetivo de Aprendizaje: ORIE05 OA 04", "Evaluación Formativa y Reflexiva Año 2026")
    ]
    saludo = (
        "Estimados apoderados y estudiantes, esperando que se encuentren bien, hago envío de los contenidos "
        "de la asignatura de Orientación para la evaluación formativa y reflexiva año 2026. A continuación, "
        "se presentan los temas centrales sobre autocuidado, toma de decisiones y prevención del consumo de drogas:"
    )
    secciones_datos = {
        "headers": ["Eje Temático", "Conceptos Fundamentales", "Habilidades y Actitudes Promovidas"],
        "widths": [Pt(145), Pt(165), Pt(201.2)],
        "filas": [
            ("Autocuidado y Prevención del Consumo de Drogas",
             "¿Qué son las drogas y sustancias nocivas para la salud? (Alcohol, tabaco, fármacos sin prescripción médica y drogas ilícitas). Efectos negativos en el desarrollo físico, cerebral, emocional y social.",
             "Reconocimiento de situaciones de riesgo y valoración de la salud y el propio cuerpo como prioridad de vida."),
            ("Factores de Riesgo vs. Factores Protectores",
             "Diferenciación entre:\n"
             "• Factores de riesgo: Presión negativa de grupos de pares, desinformación, aislamiento y curiosidad sin orientación.\n"
             "• Factores protectores: Comunicación cercana con la familia, práctica regular de deportes, amistades constructivas y metas personales.",
             "Identificación de redes de apoyo seguras en el hogar y en la comunidad escolar del Colegio Castelgandolfo."),
            ("Habilidades para la Vida: Asertividad y Decisión",
             "Capacidad de decir con firmeza y respeto 'NO' frente a situaciones peligrosas o de presión. Toma de decisiones informadas, responsables y autónomas.",
             "Desarrollo de la autoestima, pensamiento crítico, empatía y búsqueda oportuna de ayuda con adultos significativos."),
            ("Compromiso Personal y Vida Saludable",
             "Propuestas y compromisos individuales para promover un estilo de vida saludable, activo y protector dentro y fuera del colegio.",
             "Elaboración de reflexiones personales y mensajes preventivos dirigidos a los pares.")
        ]
    }
    orientaciones = [
        ("Diálogo abierto en el hogar", "Conversar en familia sobre los peligros del consumo de drogas y la importancia de la confianza mutua para plantear dudas sin temor."),
        ("Práctica de respuestas asertivas", "Modelar con el estudiante situaciones hipotéticas donde deba decir 'NO' ante presiones de otras personas con seguridad."),
        ("Refuerzo de factores protectores", "Incentivar los intereses deportivos, artísticos y culturales del estudiante como pilares de un crecimiento pleno y saludable.")
    ]
    doc = build_temario_base(
        curso_header="5° BÁSICO A",
        subject_header="ORIENTACIÓN",
        title_banner="TEMARIO Y GUÍA DE ESTUDIO: AUTOCUIDADO Y PREVENCIÓN DE DROGAS",
        datos_caja=datos_caja,
        saludo_parrafo=saludo,
        secciones_datos=secciones_datos,
        orientaciones_estudio=orientaciones,
        nota_extra="La evaluación constará de análisis de situaciones cotidianas, preguntas de reflexión y compromiso de autocuidado."
    )
    out_name = "Temario_Orientacion_5Basico_A.docx"
    safe_save(doc, os.path.join(OUTPUT_TEMARIOS_DIR, out_name))
    safe_save(doc, os.path.join(DOWNLOADS_TEMARIOS_DIR, out_name))
    print(f"[OK] Generado: {out_name}")

def reubicar_temarios_anteriores():
    """Elimina las copias de temarios sueltas en las carpetas de pruebas para centralizarlas en 00 - Temarios por Curso"""
    old_paths = [
        r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita\01 - Ciencias Naturales\5° Básico (5°A y 5°B)\Temario_Evaluacion_Final_Ciencias_5Basico.docx",
        r"C:\Users\Jack\Downloads\Evaluaciones Finales Profesora Margarita\01 - Ciencias Naturales\6° Básico (6°A y 6°B)\Temario_Evaluacion_Final_Ciencias_6Basico.docx"
    ]
    for p in old_paths:
        if os.path.exists(p):
            try:
                os.remove(p)
                print(f"[REUBICADO] Eliminado de carpeta de evaluación: {p}")
            except Exception as e:
                print(f"[AVISO] No se pudo eliminar {p}: {e}")

def generar_todos_los_temarios():
    print("\n=======================================================")
    print(" GENERANDO TEMARIOS POR CURSO PARA APODERADOS Y ALUMNOS")
    print("=======================================================")
    reubicar_temarios_anteriores()
    generar_temario_5ciencias()
    generar_temario_6ciencias()
    generar_temario_8ciencias()
    generar_temario_1musica()
    generar_temario_2musica()
    generar_temario_5orientacion()
    print("\n[ÉXITO] Todos los temarios por curso fueron generados y desplegados.")
    print(f"Carpeta de entrega: {DOWNLOADS_TEMARIOS_DIR}\n")

if __name__ == "__main__":
    generar_todos_los_temarios()
