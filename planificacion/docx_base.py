# -*- coding: utf-8 -*-
"""Utilidades de bajo nivel para armar los .docx (estilos, bordes, celdas). Sin logica pedagogica."""
import io
import os
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LOGO_DEFECTO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'assets', 'logo_colegio.png')
NAVY, MID, ICE, PALE, LINE = '1F3864', '2E5597', 'DCE6F2', 'F4F7FB', '8EA9DB'
CENTER = WD_ALIGN_PARAGRAPH.CENTER
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


def nuevo_documento(ancho_cm=33.0, alto_cm=21.6, margen_lat=1.5, margen_sup=1.0, margen_inf=1.1):
    """Hoja oficio horizontal (formato del colegio)."""
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(ancho_cm), Cm(alto_cm)
    sec.left_margin = sec.right_margin = Cm(margen_lat)
    sec.top_margin, sec.bottom_margin = Cm(margen_sup), Cm(margen_inf)
    sec.footer_distance = Cm(0.5)
    return doc


def run(p, text, bold=False, size=10, color=None, italic=False):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    r.font.name = 'Arial'; r.font.size = Pt(size)
    rf = r._r.get_or_add_rPr().find(qn('w:rFonts'))
    rf.set(qn('w:eastAsia'), 'Arial'); rf.set(qn('w:cs'), 'Arial')
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def parrafo(cell, text='', bold=False, size=10, color=None, align=None, italic=False, first=True,
            after=0, left=None, hanging=None):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after); pf.space_before = Pt(0); pf.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    if left is not None:
        pf.left_indent = Cm(left / 567); pf.first_line_indent = Cm(-(hanging or 0) / 567)
    if text:
        run(p, text, bold, size, color, italic)
    return p


def sombrear(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for s in tcPr.findall(qn('w:shd')):
        tcPr.remove(s)
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    pos = 0
    for i, ch in enumerate(list(tcPr)):
        if ch.tag in (qn('w:tcW'), qn('w:gridSpan'), qn('w:vMerge'), qn('w:tcBorders')):
            pos = i + 1
    tcPr.insert(pos, sh)


def borde_celda(cell, color=MID, sz='6'):
    pr = cell._tc.get_or_add_tcPr()
    for b in pr.findall(qn('w:tcBorders')):
        pr.remove(b)
    tb = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        tb.append(e)
    pos = 0
    for i, ch in enumerate(list(pr)):
        if ch.tag in (qn('w:tcW'), qn('w:gridSpan'), qn('w:vMerge')):
            pos = i + 1
    pr.insert(pos, tb)


def valign(cell, v='center'):
    e = OxmlElement('w:vAlign'); e.set(qn('w:val'), v); cell._tc.get_or_add_tcPr().append(e)


def tabla(doc, filas, anchos, borde=LINE, margenes=(60, 100)):
    """Tabla de ancho fijo (twips) con bordes finos y margenes de celda."""
    t = doc.add_table(rows=filas, cols=len(anchos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = t._tbl; tblPr = tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = tblPr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tblPr.append(tw)
    tw.set(qn('w:w'), str(sum(anchos))); tw.set(qn('w:type'), 'dxa')
    for g in list(tbl.tblGrid):
        tbl.tblGrid.remove(g)
    for w in anchos:
        g = OxmlElement('w:gridCol'); g.set(qn('w:w'), str(w)); tbl.tblGrid.append(g)
    for row in t.rows:
        for c, w in zip(row.cells, anchos):
            c.width = Cm(w / 567)
    if borde:
        bd = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{side}')
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), borde)
            bd.append(e)
        tblPr.append(bd)
    mar = OxmlElement('w:tblCellMar')
    for k, v in (('top', margenes[0]), ('left', margenes[1]), ('bottom', margenes[0]), ('right', margenes[1])):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)
    return t


def fila_fija(row, encabezado=False, alto_min=None):
    """La fila no se parte entre paginas; el encabezado se repite; alto minimo opcional (twips)."""
    trPr = row._tr.get_or_add_trPr()
    if alto_min:
        h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(alto_min)); h.set(qn('w:hRule'), 'atLeast'); trPr.append(h)
    trPr.append(OxmlElement('w:cantSplit'))
    if encabezado:
        trPr.append(OxmlElement('w:tblHeader'))


def espaciador(doc, pt=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(pt)
    r = p.add_run(''); r.font.size = Pt(2)
    return p


def pie_de_pagina(doc, texto):
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = CENTER
    run(fp, texto + '  ·  Página ', False, 8, '595959')
    r = fp.add_run(); r.font.size = Pt(8); r.font.name = 'Arial'
    for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if typ:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), typ); r._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(XML_SPACE, 'preserve'); it.text = txt; r._r.append(it)


def franja_titulo(doc, logo_path, institucion, titulo, subtitulo, ancho, salto_pagina=False):
    """Logo + nombre del colegio + titulo en una sola franja (igual en todos los documentos)."""
    t = tabla(doc, 1, [1000, 3700, ancho - 4700], borde=None)
    a, b, c = t.rows[0].cells
    p = a.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    if logo_path and os.path.exists(logo_path):
        p.add_run().add_picture(logo_path, height=Cm(1.25))
    lineas = list(institucion)
    for i, (txt, sz, col, bold) in enumerate(lineas):
        pp = b.paragraphs[0] if i == 0 else b.add_paragraph()
        pp.paragraph_format.space_after = Pt(0)
        run(pp, txt, bold, sz, col)
    sombrear(c, NAVY)
    parrafo(c, titulo, True, 14, 'FFFFFF', CENTER)
    parrafo(c, subtitulo, True, 10, ICE, CENTER, first=False)
    for cell in (a, b, c):
        valign(cell)
    if salto_pagina:
        a.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
    espaciador(doc, 5)
    return t
