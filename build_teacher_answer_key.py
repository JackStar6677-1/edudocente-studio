import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from generator_core import (
    create_base_doc, add_header, add_title_banner, add_student_info,
    add_section_header, set_cell_shading, set_cell_borders, set_cell_margins,
    COLOR_NAVY_HEX, COLOR_ICE_HEX, COLOR_BOX_BG_HEX, COLOR_BORDER_HEX,
    COLOR_CORRECT_HEX, COLOR_NAVY_RGB, COLOR_CORRECT_RGB, COLOR_DARK_RGB,
    COLOR_WHITE_RGB, ASSETS_DIR, OUTPUT_DIR
)

def build_pauta_correccion():
    doc = create_base_doc()

    # ================== PÁGINA 1 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 5° BÁSICO A - B", subtitle="DOCUMENTO DOCENTE")
    add_title_banner(doc, "PAUTA DE CORRECCIÓN: EVALUACIÓN FINAL DE CIENCIAS NATURALES", is_pauta=True)
    add_student_info(doc, is_pauta=True)

    # SOLUCIONARIO ÍTEM I
    add_section_header(doc, "SOLUCIONARIO ÍTEM I: SELECCIÓN MÚLTIPLE (2 pts c/u — Total: 10 pts)")

    pautas_item1 = [
        ("1. ¿Qué es la energía eléctrica?",
         "C) Una forma de energía que se produce por el movimiento y flujo de cargas eléctricas.",
         "Justificación: La corriente eléctrica corresponde al movimiento ordenado de cargas eléctricas a través de un conductor. Texto escolar Pág. 146."),

        ("2. Esta energía surge de las propiedades de la materia. ¿Cómo se describe la materia?",
         "A) La materia está formada por partículas diminutas que poseen cargas eléctricas.",
         "Justificación: Toda la materia está compuesta por átomos con partículas cargadas positiva (protones) y negativamente (electrones). Texto escolar Pág. 146."),

        ("3. En los modelos científicos, las partículas que componen la materia se suelen representar como:",
         "B) Pequeñas esferas.",
         "Justificación: En el modelo corpuscular escolar, las partículas atómicas se ilustran convencionalmente como esferas diminutas. Texto escolar Pág. 146."),

        ("4. Las partículas pueden adquirir o manifestar una carga eléctrica, la que puede ser:",
         "B) Positiva (+) o negativa (-).",
         "Justificación: Existen dos naturalezas fundamentales de carga eléctrica: positiva (+) y negativa (-). Texto escolar Pág. 146."),

        ("5. Materiales por donde se desplazan con facilidad las cargas negativas:",
         "B) Cobre.",
         "Justificación: El cobre es un metal de excelente conductividad eléctrica. Plástico, madera y goma son aislantes. Texto escolar Pág. 152.")
    ]

    for q_t, ans, just in pautas_item1:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(2)
        p_q.paragraph_format.space_after = Pt(1)
        rq = p_q.add_run(q_t)
        rq.bold = True
        rq.font.name = "Arial"
        rq.font.size = Pt(8.5)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Inches(0.2)
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(0.5)
        ra_b = p_ans.add_run("Respuesta Correcta: ")
        ra_b.bold = True
        ra_b.font.name = "Arial"
        ra_b.font.size = Pt(8.5)
        ra_b.font.color.rgb = COLOR_CORRECT_RGB
        ra_t = p_ans.add_run(ans)
        ra_t.bold = True
        ra_t.font.name = "Arial"
        ra_t.font.size = Pt(8.5)
        ra_t.font.color.rgb = COLOR_CORRECT_RGB

        p_just = doc.add_paragraph()
        p_just.paragraph_format.left_indent = Inches(0.2)
        p_just.paragraph_format.space_before = Pt(0)
        p_just.paragraph_format.space_after = Pt(2)
        rj = p_just.add_run(just)
        rj.font.name = "Arial"
        rj.font.size = Pt(8)
        rj.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_page_break()

    # ================== PÁGINA 2 ==================
    add_header(doc, school="COLEGIO CASTELGANDOLFO", subject="CIENCIAS NATURALES - 5° BÁSICO A - B", subtitle="PAUTA DE CORRECCIÓN (PÁG. 2)")

    # SOLUCIONARIO ÍTEM II
    add_section_header(doc, "SOLUCIONARIO ÍTEM II: VERDADERO O FALSO (1 pto c/u — Total: 10 pts)")

    pautas_vf = [
        ("V", "1. Los aparatos electrónicos y electrodomésticos funcionan gracias a los circuitos eléctricos.",
         "Verdadero: Todo artefacto eléctrico requiere un circuito para conducir la corriente hacia sus partes operativas."),
        
        ("V", "2. Un circuito eléctrico es un camino cerrado por el que fluyen y circulan las cargas eléctricas.",
         "Verdadero: Para que haya flujo continuo de cargas, el lazo conductor debe ser cerrado."),
        
        ("F", "3. Todos los circuitos eléctricos tienen componentes que funcionan sin conectarse entre sí.",
         "Falsa: Si los componentes están desconectados, el circuito queda abierto y no fluye corriente eléctrica."),
        
        ("V", "4. El generador o fuente de energía proporciona la energía necesaria para que circule la corriente eléctrica.",
         "Verdadero: Proporciona la fuerza impulsora (voltaje) indispensable para el movimiento de cargas."),
        
        ("F", "5. Los cables conducen corriente eléctrica por las paredes de forma espontánea sin fuente de energía.",
         "Falsa: Los cables son conductores pasivos; sin una fuente que impulse las cargas, no circula corriente."),
        
        ("V", "6. La resistencia o receptor recibe y transforma la energía eléctrica en otro tipo de energía.",
         "Verdadero: Transforma la electricidad en energía útil (lumínica, térmica, cinética o sonora)."),
        
        ("V", "7. Si el receptor es una ampolleta, la energía eléctrica se transforma principalmente en energía lumínica.",
         "Verdadero: La función de la ampolleta es emitir luz visible (acompañada secundariamente de calor)."),
        
        ("F", "8. El interruptor regula el paso de la corriente hidráulica.",
         "Falsa: Regula el paso de corriente ELÉCTRICA (abriendo o cerrando el circuito eléctrico)."),
        
        ("V", "9. Cuando el interruptor está abierto, la corriente eléctrica deja de circular por el circuito.",
         "Verdadero: Al estar abierto se interrumpe el camino conductor e impide el paso de electrones."),
        
        ("V", "10. Cuando el interruptor está cerrado, la corriente eléctrica circula con normalidad.",
         "Verdadero: Los contactos metálicos se unen, completando el camino conductor.")
    ]

    for v_f, sent, just in pautas_vf:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)

        badge = f"[  {v_f}  ]  "
        rb = p.add_run(badge)
        rb.bold = True
        rb.font.name = "Arial"
        rb.font.size = Pt(8.5)
        if v_f == "V":
            rb.font.color.rgb = COLOR_CORRECT_RGB
        else:
            rb.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

        rs = p.add_run(sent)
        rs.bold = True
        rs.font.name = "Arial"
        rs.font.size = Pt(8.5)

        pj = doc.add_paragraph()
        pj.paragraph_format.left_indent = Inches(0.4)
        pj.paragraph_format.space_before = Pt(0)
        pj.paragraph_format.space_after = Pt(2.5)
        rj = pj.add_run(f"→ {just}")
        rj.font.name = "Arial"
        rj.font.size = Pt(8)
        rj.italic = True
        rj.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # SOLUCIONARIO ÍTEM III
    add_section_header(doc, "SOLUCIONARIO ÍTEM III: SÍMBOLOS SEGÚN CORRESPONDA (2 pts c/u — Total: 8 pts)")

    t_cuadros = doc.add_table(rows=2, cols=2)
    t_cuadros.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cuadros.autofit = False

    comp_pauta = [
        {"title": "1. Fuente o generador", "img_sym": "simbolo_pila.png",
         "criterio": "2 pts: Placas paralelas desiguales: línea larga (+) y línea corta/gruesa con signo menos (–).", "row": 0, "col": 0},
        {"title": "2. Cables", "img_sym": "simbolo_cable.png",
         "criterio": "2 pts: Línea continua, recta y limpia que representa el cable metálico conductor.", "row": 0, "col": 1},
        {"title": "3. Interruptor", "img_sym": "simbolo_interruptor.png",
         "criterio": "2 pts: Dos bornes con palanca levantada (abierto) o palanca unida (cerrado). Cualquiera es válido.", "row": 1, "col": 0},
        {"title": "4. Resistencia", "img_sym": "simbolo_ampolleta.png",
         "criterio": "2 pts: Círculo al centro con 'X' interior y dos líneas laterales (ampolleta/receptor) o símbolo de resistencia (zigzag/rectángulo).", "row": 1, "col": 1},
    ]

    for comp in comp_pauta:
        cell = t_cuadros.cell(comp["row"], comp["col"])
        cell.width = Pt(250)
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        set_cell_borders(cell, top={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               bottom={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               left={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX},
                               right={'val': 'single', 'sz': '8', 'color': COLOR_BORDER_HEX})
        set_cell_shading(cell, "F6FAFE")

        p_h = cell.paragraphs[0]
        p_h.paragraph_format.space_before = Pt(0)
        p_h.paragraph_format.space_after = Pt(2)
        rh = p_h.add_run(comp['title'])
        rh.bold = True
        rh.font.name = "Arial"
        rh.font.size = Pt(9.5)
        rh.font.color.rgb = COLOR_NAVY_RGB

        # Recuadro con el símbolo solución oficial
        t_sym = cell.add_table(rows=1, cols=1)
        t_sym.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_sym.autofit = False
        c_sym = t_sym.cell(0, 0)
        c_sym.width = Pt(230)
        set_cell_margins(c_sym, top=30, bottom=30, left=30, right=30)
        set_cell_borders(c_sym, top={'val': 'single', 'sz': '8', 'color': COLOR_CORRECT_HEX},
                                bottom={'val': 'single', 'sz': '8', 'color': COLOR_CORRECT_HEX},
                                left={'val': 'single', 'sz': '8', 'color': COLOR_CORRECT_HEX},
                                right={'val': 'single', 'sz': '8', 'color': COLOR_CORRECT_HEX})
        set_cell_shading(c_sym, "FFFFFF")

        p_is = c_sym.paragraphs[0]
        p_is.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_is.paragraph_format.space_before = Pt(0)
        p_is.paragraph_format.space_after = Pt(0)

        f_sym = os.path.join(ASSETS_DIR, comp["img_sym"])
        if os.path.exists(f_sym):
            p_is.add_run().add_picture(f_sym, width=Inches(1.8))

        # Criterio
        p_crit = cell.add_paragraph()
        p_crit.paragraph_format.space_before = Pt(3)
        p_crit.paragraph_format.space_after = Pt(1)
        rc = p_crit.add_run(comp["criterio"])
        rc.font.name = "Arial"
        rc.font.size = Pt(7.5)
        rc.italic = True
        rc.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    out_path = os.path.join(OUTPUT_DIR, "Pauta_Correccion_Ciencias_5Basico.docx")
    doc.save(out_path)
    print(f"Pauta actualizada con éxito en: {out_path}")

if __name__ == "__main__":
    build_pauta_correccion()
